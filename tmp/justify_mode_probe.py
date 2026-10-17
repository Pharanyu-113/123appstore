from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


FONT = "TH Sarabun PSK"
TEXT = "ปัจจุบันเทคโนโลยีดิจิทัลมีบทบาทต่อการดำเนินธุรกิจทุกระดับ โดยเฉพาะร้านค้าปลีกส่งที่ต้องบริหารข้อมูลสินค้า คำสั่งซื้อ และการจัดส่งให้ทันต่อการใช้งานจริง"


def set_font(run):
    run.font.name = FONT
    run.font.size = Pt(16)
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts or OxmlElement("w:rFonts")
    if fonts.getparent() is None:
        rpr.insert(0, fonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attr}"), FONT)


doc = Document()
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.left_margin = Cm(3.81)
section.right_margin = Cm(2.54)
section.top_margin = Cm(2.54)
section.bottom_margin = Cm(2.54)

for label, mode in (
    ("JUSTIFY", WD_ALIGN_PARAGRAPH.JUSTIFY),
    ("THAI_JUSTIFY", WD_ALIGN_PARAGRAPH.THAI_JUSTIFY),
    ("JUSTIFY_LOW", WD_ALIGN_PARAGRAPH.JUSTIFY_LOW),
    ("JUSTIFY_MED", WD_ALIGN_PARAGRAPH.JUSTIFY_MED),
    ("JUSTIFY_HI", WD_ALIGN_PARAGRAPH.JUSTIFY_HI),
    ("DISTRIBUTE", WD_ALIGN_PARAGRAPH.DISTRIBUTE),
    ("LEFT", WD_ALIGN_PARAGRAPH.LEFT),
):
    h = doc.add_paragraph()
    h.add_run(label).bold = True
    for run in h.runs:
        set_font(run)
    p = doc.add_paragraph()
    p.alignment = mode
    p.paragraph_format.first_line_indent = Cm(1.27)
    p.add_run(TEXT)
    for run in p.runs:
        set_font(run)

for label, text in (
    ("JUSTIFY_WITH_NBSP", TEXT.replace(" ", "\u00A0")),
    ("JUSTIFY_WITH_ZWSP", TEXT.replace(" ", "\u200B")),
    ("JUSTIFY_WITHOUT_SPACES", TEXT.replace(" ", "")),
):
    h = doc.add_paragraph()
    h.add_run(label).bold = True
    for run in h.runs:
        set_font(run)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(1.27)
    p.add_run(text)
    for run in p.runs:
        set_font(run)

doc.save("tmp/justify_mode_probe.docx")
