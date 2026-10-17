"""Render-only experiment for Word justification with Thai and English terms.

This script deliberately lives in tmp/ and never opens or modifies the report.
"""

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


FONT = "TH Sarabun PSK"
THAI = r"\u0E00-\u0E7F"
THAI_GAP = re.compile(rf"(?<=[{THAI}]) (?=[{THAI}])")
SAMPLE = (
    "การพัฒนาแบบ Full Stack คือการพัฒนาที่ครอบคลุมทั้งส่วนติดต่อผู้ใช้และส่วนทำงานเบื้องหลังของระบบ "
    "ในโครงงานนี้ส่วนหน้าบ้านใช้ Next.js และ React เพื่อแสดงหน้าจอให้เหมาะกับบทบาทของผู้ใช้ "
    "ส่วนหลังบ้านใช้ API route ของ Next.js สำหรับตรวจสอบ session รับข้อมูลและเรียกข้อมูลจาก Supabase PostgreSQL "
    "การออกแบบลักษณะนี้ช่วยให้ส่วนที่ผู้ใช้มองเห็นและส่วนที่จัดการข้อมูลทำงานเชื่อมกันอย่างเป็นระบบ"
)


def transform(text: str, visible_gap: str) -> str:
    """Use a zero-width opportunity between Thai words; vary every other gap."""
    text = THAI_GAP.sub("\u200B", text)
    return text.replace(" ", visible_gap)


def set_font(run):
    run.font.name = FONT
    run.font.size = Pt(16)
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts or OxmlElement("w:rFonts")
    if fonts.getparent() is None:
        rpr.insert(0, fonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attr}"), FONT)


def add_text(doc, text, *, heading=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if heading else WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(0) if heading else Cm(1.27)
    p.paragraph_format.space_before = Pt(3 if heading else 0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(text)
    if heading:
        run.bold = True
    set_font(run)


def main():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3.81)
    section.right_margin = Cm(2.54)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)

    variants = (
        ("A. standard ASCII space (baseline)", " "),
        ("B. NBSP U+00A0", "\u00A0"),
        ("C. narrow no-break space U+202F", "\u202F"),
        ("D. thin space U+2009", "\u2009"),
        ("E. figure space U+2007", "\u2007"),
        ("F. word-joiner + thin space + word-joiner", "\u2060\u2009\u2060"),
        ("G. word-joiner + narrow NBSP + word-joiner", "\u2060\u202F\u2060"),
    )
    add_text(doc, "Word justification probe — Thai zero-width breaks + English phrase gaps", heading=True)
    for label, gap in variants:
        add_text(doc, label, heading=True)
        add_text(doc, transform(SAMPLE, gap))

    out = Path("tmp/english_spacing_probe.docx")
    doc.save(out)
    print(out)


if __name__ == "__main__":
    main()
