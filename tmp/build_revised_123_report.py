from pathlib import Path
from datetime import date

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(r"C:\store-app\123app")
OUT_DOCX = ROOT / "output" / "รายงานโครงงาน_123พาณิชย์ปลีกส่ง_ฉบับปรับปรุงขอบเขตและแผนงาน_2569.docx"
ASSETS = ROOT / "tmp" / "report_assets_revised"
FONT = "TH Sarabun PSK"
FONT_FILE = r"C:\store-app\123app\tmp\report_fonts\THSarabunPSK-Regular.ttf"
BOLD_FONT_FILE = r"C:\store-app\123app\tmp\report_fonts\THSarabunPSK-Bold.ttf"
INK = "1F2933"
GREEN = "0F5D52"
PALE = "EAF3F1"
LINE = "91ACA6"


def set_run_font(run, size=16, bold=None, color=INK, font=FONT):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:ascii"), font)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), font)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run._element.rPr.rFonts.set(qn("w:cs"), font)
    run._element.rPr.rFonts.set(qn("w:hint"), "eastAsia")
    lang = run._element.rPr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        run._element.rPr.append(lang)
    lang.set(qn("w:val"), "th-TH")
    lang.set(qn("w:eastAsia"), "th-TH")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return run


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=LINE, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_table_border(table, color="000000", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def set_paragraph_keep(paragraph, keep_next=False):
    ppr = paragraph._p.get_or_add_pPr()
    element = OxmlElement("w:keepNext" if keep_next else "w:widowControl")
    element.set(qn("w:val"), "true")
    ppr.append(element)


def add_field(paragraph, field):
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(sep)
    run._r.append(text)
    run._r.append(end)
    set_run_font(run, 13, color=INK)


def set_page_number_format(section, fmt="decimal", start=None):
    sect_pr = section._sectPr
    node = sect_pr.find(qn("w:pgNumType"))
    if node is None:
        node = OxmlElement("w:pgNumType")
        sect_pr.append(node)
    node.set(qn("w:fmt"), fmt)
    if start is not None:
        node.set(qn("w:start"), str(start))


def add_footer_number(section, center=False):
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    add_field(p, "PAGE")


def configure_section(section, footer_center=False, number_fmt="decimal", start=None, show_page_number=True):
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(25.4)
    section.bottom_margin = Mm(22.0)
    section.left_margin = Mm(31.75)
    section.right_margin = Mm(25.4)
    section.header_distance = Mm(10)
    section.footer_distance = Mm(11)
    if show_page_number:
        set_page_number_format(section, number_fmt, start)
        add_footer_number(section, footer_center)


def add_text(doc, text, size=16, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT,
             space_before=0, space_after=4, first_indent=True, line_spacing=1.18):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line_spacing
    if first_indent:
        pf.first_line_indent = Mm(12.7)
    for index, chunk in enumerate(text.split("\n")):
        if index:
            p.add_run().add_break()
        set_run_font(p.add_run(chunk), size=size, bold=bold)
    return p


def add_heading(doc, text, level=1, number=None, page_break_before=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if level > 0 else WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.space_before = Pt(11 if level == 1 else 7)
    pf.space_after = Pt(5)
    pf.line_spacing = 1.15
    pf.page_break_before = page_break_before
    r = p.add_run((number + " " if number else "") + text)
    set_run_font(r, 22 if level == 1 else 18, bold=True, color="000000")
    set_paragraph_keep(p, keep_next=True)
    return p


def add_caption(doc, text, figure=True):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if figure else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(7)
    set_run_font(p.add_run(text), 13, bold=True)
    set_paragraph_keep(p, keep_next=True)
    return p


def add_figure(doc, image_path, caption, width=Mm(155)):
    """Keep a diagram and its caption together when Word paginates the report."""
    doc.add_picture(str(image_path), width=width)
    image_paragraph = doc.paragraphs[-1]
    set_paragraph_keep(image_paragraph, keep_next=True)
    add_caption(doc, caption)


def add_table(doc, caption, headers, rows, widths=None, font_size=13):
    font_size = max(font_size, 13)
    add_caption(doc, caption, figure=False)
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    head = table.rows[0]
    set_repeat_table_header(head)
    for idx, value in enumerate(headers):
        cell = head.cells[idx]
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(cell, PALE)
        set_cell_border(cell)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_run_font(p.add_run(str(value)), font_size, bold=True, color="000000")
        if widths:
            cell.width = widths[idx]
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cell = cells[idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_border(cell)
            set_cell_margins(cell)
            if widths:
                cell.width = widths[idx]
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.18
            for line_index, piece in enumerate(str(value).split("\n")):
                if line_index:
                    p.add_run().add_break()
                set_run_font(p.add_run(piece), font_size)
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_scope_group(doc, number, title, bullets):
    """Add a formal numbered scope group with hanging hyphen bullets."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.left_indent = Mm(12.7)
    p.paragraph_format.first_line_indent = Mm(-6.35)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.18
    set_run_font(p.add_run(f"{number}) "), 16, bold=True, color="000000")
    set_run_font(p.add_run(title), 16, bold=True, color="000000")
    set_paragraph_keep(p, keep_next=True)
    for item in bullets:
        bullet = doc.add_paragraph()
        bullet.alignment = WD_ALIGN_PARAGRAPH.LEFT
        bullet.paragraph_format.left_indent = Mm(20)
        bullet.paragraph_format.first_line_indent = Mm(-5)
        bullet.paragraph_format.space_after = Pt(1)
        bullet.paragraph_format.line_spacing = 1.18
        set_run_font(bullet.add_run("- "), 16)
        set_run_font(bullet.add_run(item), 16)


def add_schedule_table(doc):
    """Create a compact Gantt-style project-plan table matching the reference layout."""
    add_caption(doc, "ตารางที่ 1.3 ระยะเวลาการดำเนินงาน", figure=False)
    months = ["ก.พ.", "มี.ค.", "เม.ย.", "พ.ค.", "มิ.ย.", "ก.ค.", "ส.ค.", "ก.ย.", "ต.ค."]
    activities = [
        ("1. ศึกษาและรวบรวมข้อมูล", 0, 0),
        ("2. วิเคราะห์และออกแบบระบบ", 1, 2),
        ("3. ออกแบบหน้าจอ", 2, 2),
        ("4. พัฒนาระบบ", 2, 3),
        ("5. ทดสอบระบบและแก้ไข", 3, 4),
        ("6. ประเมินประสิทธิภาพระบบโดยผู้เชี่ยวชาญ", 4, 5),
        ("7. ทำเอกสาร", 6, 7),
        ("8. วิเคราะห์ข้อมูลและสรุปผลจัดทำรายงาน", 7, 8),
    ]
    table = doc.add_table(rows=2, cols=10)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    activity_width = Mm(51)
    month_width = Mm(11.2)
    for row in table.rows:
        row.cells[0].width = activity_width
        for cell in row.cells[1:]:
            cell.width = month_width
    activity_head = table.cell(0, 0).merge(table.cell(1, 0))
    duration_head = table.cell(0, 1).merge(table.cell(0, 9))
    for cell, text in [(activity_head, "กิจกรรม"), (duration_head, "ระยะเวลาในการดำเนินงาน (ปี พ.ศ. 2569)")]:
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_border(cell, color="000000", size="6")
        set_cell_margins(cell, top=70, start=70, bottom=70, end=70)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_run_font(p.add_run(text), 12, bold=True, color="000000")
    for idx, month in enumerate(months, 1):
        cell = table.cell(1, idx)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_border(cell, color="000000", size="6")
        set_cell_margins(cell, top=55, start=22, bottom=55, end=22)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_run_font(p.add_run(month), 11.5, bold=True, color="000000")
    set_repeat_table_header(table.rows[0])
    set_repeat_table_header(table.rows[1])
    for activity, start, end in activities:
        row = table.add_row()
        row.height = Mm(14)
        activity_cell = row.cells[0]
        activity_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_border(activity_cell, color="000000", size="6")
        set_cell_margins(activity_cell, top=65, start=90, bottom=65, end=80)
        p = activity_cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.08
        set_run_font(p.add_run(activity), 12, color="000000")
        for idx in range(1, 10):
            cell = row.cells[idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_border(cell, color="000000", size="6")
            set_cell_margins(cell, top=45, start=22, bottom=45, end=22)
        active = row.cells[start + 1]
        if end > start:
            active = active.merge(row.cells[end + 1])
        active.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_border(active, color="000000", size="6")
        set_cell_margins(active, top=45, start=22, bottom=45, end=22)
        p = active.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run()
        span = (end - start + 1) * 11.2 - 0.8
        run.add_picture(str(ASSETS / "schedule-arrow.png"), width=Mm(span), height=Mm(2.1))
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
    set_table_border(table, color="000000", size="6")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def image_font(size, bold=False):
    return ImageFont.truetype(BOLD_FONT_FILE if bold else FONT_FILE, size=size)


def rounded(draw, box, fill, outline=GREEN, radius=18, width=3):
    if not fill.startswith("#"):
        fill = "#" + fill
    if not outline.startswith("#"):
        outline = "#" + outline
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def center(draw, box, text, font, fill=INK):
    if not fill.startswith("#"):
        fill = "#" + fill
    left, top, right, bottom = box
    bbox = draw.multiline_textbbox((0, 0), text, font=font, align="center", spacing=4)
    x = (left + right - (bbox[2] - bbox[0])) / 2
    y = (top + bottom - (bbox[3] - bbox[1])) / 2
    draw.multiline_text((x, y), text, font=font, fill=fill, align="center", spacing=4)


def arrow(draw, start, end, color=GREEN, width=5):
    if not color.startswith("#"):
        color = "#" + color
    draw.line([start, end], fill=color, width=width)
    x1, y1 = end
    x0, y0 = start
    if abs(x1 - x0) >= abs(y1 - y0):
        sign = 1 if x1 >= x0 else -1
        poly = [(x1, y1), (x1 - 16 * sign, y1 - 10), (x1 - 16 * sign, y1 + 10)]
    else:
        sign = 1 if y1 >= y0 else -1
        poly = [(x1, y1), (x1 - 10, y1 - 16 * sign), (x1 + 10, y1 - 16 * sign)]
    draw.polygon(poly, fill=color)


def draw_architecture(path):
    im = Image.new("RGB", (1800, 900), "white")
    d = ImageDraw.Draw(im)
    title = image_font(50, True)
    h = image_font(36, True)
    body = image_font(29)
    center(d, (0, 25, 1800, 90), "SYSTEM ARCHITECTURE", title, GREEN)
    items = [
        ((70, 220, 390, 600), "USERS", "OWNER\nADMIN\nEMPLOYEE\nRETAILER"),
        ((530, 220, 910, 600), "WEB CLIENT", "React components\nCart & GPS\nRole dashboards"),
        ((1050, 180, 1430, 640), "NEXT.JS SERVER", "Route handlers\nSession & RBAC\nValidation"),
        ((1480, 220, 1740, 600), "SUPABASE", "PostgreSQL\nREST / RPC\nRLS"),
    ]
    for box, heading, body_text in items:
        rounded(d, box, "#EAF3F1")
        center(d, (box[0]+15, box[1]+35, box[2]-15, box[1]+105), heading, h, GREEN)
        center(d, (box[0]+15, box[1]+135, box[2]-15, box[3]-35), body_text, body)
    for a, b in [((390, 410), (530, 410)), ((910, 410), (1050, 410)), ((1430, 410), (1480, 410))]:
        arrow(d, a, b)
    rounded(d, (1050, 700, 1740, 825), "#F6FAF9", LINE, 12, 2)
    center(d, (1070, 720, 1720, 805), "OpenStreetMap embed  |  Google Maps link", body, INK)
    arrow(d, (1240, 640), (1240, 700), LINE, 4)
    im.save(path)


def draw_order_flow(path):
    im = Image.new("RGB", (1800, 660), "white")
    d = ImageDraw.Draw(im)
    title = image_font(50, True)
    h = image_font(30, True)
    body = image_font(25)
    center(d, (0, 15, 1800, 80), "ORDER AND DELIVERY FLOW", title, GREEN)
    steps = [
        (70, "RETAILER", "selects items\nbox / pack"),
        (350, "CHECKOUT", "validate stock\nand price"),
        (630, "ATOMIC SAVE", "order + receipt\nstock deduction"),
        (910, "PENDING", "employee\nreceives job"),
        (1190, "DELIVERING", "GPS position\nsharing"),
        (1470, "COMPLETED", "delivery\nrecorded"),
    ]
    for x, heading, detail in steps:
        box = (x, 230, x + 230, 475)
        rounded(d, box, "#EAF3F1")
        center(d, (x+10, 265, x+220, 330), heading, h, GREEN)
        center(d, (x+14, 350, x+216, 445), detail, body)
    for x, _, _ in steps[:-1]:
        arrow(d, (x+230, 352), (x+280, 352))
    im.save(path)


def draw_data_model(path):
    im = Image.new("RGB", (1800, 980), "white")
    d = ImageDraw.Draw(im)
    title = image_font(50, True)
    h = image_font(31, True)
    body = image_font(24)
    center(d, (0, 20, 1800, 90), "CORE DATA MODEL", title, GREEN)
    items = [
        ((70, 160, 360, 360), "profiles", "users, role,\nshop details"),
        ((510, 160, 810, 360), "products", "price, stock,\ncategory"),
        ((960, 160, 1250, 360), "orders", "retailer, employee,\nstatus"),
        ((1400, 160, 1730, 360), "daily_receipts", "one retailer\nper day"),
        ((520, 590, 810, 800), "order_items", "quantity, price,\nhistorical name"),
        ((990, 590, 1280, 800), "receipt_items", "daily receipt\nline items"),
        ((1420, 590, 1720, 800), "delivery_tracking", "last GPS\nposition"),
    ]
    for box, heading, detail in items:
        rounded(d, box, "#F6FAF9", LINE, 16, 3)
        center(d, (box[0]+10, box[1]+25, box[2]-10, box[1]+90), heading, h, GREEN)
        center(d, (box[0]+10, box[1]+105, box[2]-10, box[3]-25), detail, body)
    connections = [
        ((360, 260), (510, 260)), ((810, 260), (960, 260)), ((1250, 260), (1400, 260)),
        ((665, 360), (665, 590)), ((1105, 360), (1105, 590)), ((1250, 320), (1420, 690)),
    ]
    for a, b in connections:
        arrow(d, a, b, LINE, 3)
    im.save(path)


def draw_roles(path):
    im = Image.new("RGB", (1800, 720), "white")
    d = ImageDraw.Draw(im)
    title = image_font(50, True)
    h = image_font(34, True)
    body = image_font(26)
    center(d, (0, 15, 1800, 85), "ROLE-BASED DASHBOARD ACCESS", title, GREEN)
    roles = [
        (60, "OWNER", "Catalog & stock\nSales & delivery view"),
        (480, "ADMIN", "Inventory\nadjustment"),
        (900, "EMPLOYEE", "Order status\nGPS sharing"),
        (1320, "RETAILER", "Catalog, cart, receipt\nOrder tracking"),
    ]
    for x, heading, detail in roles:
        rounded(d, (x, 230, x+340, 520), "#EAF3F1", GREEN, 20, 3)
        center(d, (x+10, 275, x+330, 345), heading, h, GREEN)
        center(d, (x+15, 375, x+325, 490), detail, body)
    im.save(path)


def draw_schedule_arrow(path):
    im = Image.new("RGB", (1200, 100), "white")
    d = ImageDraw.Draw(im)
    y = 50
    d.line([(58, y), (1142, y)], fill="black", width=12)
    d.polygon([(34, y), (92, y - 28), (92, y + 28)], fill="black")
    d.polygon([(1166, y), (1108, y - 28), (1108, y + 28)], fill="black")
    im.save(path)


def make_assets():
    ASSETS.mkdir(parents=True, exist_ok=True)
    draw_architecture(ASSETS / "architecture.png")
    draw_order_flow(ASSETS / "order-flow.png")
    draw_data_model(ASSETS / "data-model.png")
    draw_roles(ASSETS / "roles.png")
    draw_schedule_arrow(ASSETS / "schedule-arrow.png")


def cover(doc):
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    set_run_font(p.add_run("รายงานโครงงานคอมพิวเตอร์"), 22, bold=True, color="000000")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(7)
    set_run_font(p.add_run("การพัฒนาเว็บแอปพลิเคชันระบบจัดการค้าส่ง"), 26, bold=True, color="000000")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    set_run_font(p.add_run("123พาณิชย์ปลีกส่ง"), 30, bold=True, color=GREEN)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(36)
    set_run_font(p.add_run("Development of a Web Application for\n123 Phanit Retail and Wholesale Management"), 18, bold=True, color="000000")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    set_run_font(p.add_run("ภรัณยู โยธาพล"), 20, bold=True, color="000000")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(38)
    set_run_font(p.add_run("66102122113"), 18, bold=True, color="000000")
    for text, size in [
        ("โครงงานคอมพิวเตอร์นี้เป็นส่วนหนึ่งของการศึกษา", 18),
        ("ตามหลักสูตรปริญญาวิทยาศาสตรบัณฑิต", 18),
        ("สาขาวิชาเทคโนโลยีคอมพิวเตอร์และดิจิทัล", 18),
        ("มหาวิทยาลัยราชภัฏสกลนคร", 18),
        ("พ.ศ. 2569", 20),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(4)
        set_run_font(p.add_run(text), size, bold=True, color="000000")


def abstract(doc, english=False):
    title = "ABSTRACT" if english else "บทคัดย่อ"
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(16)
    set_run_font(p.add_run(title), 22, bold=True, color="000000")
    if not english:
        paragraphs = [
            "โครงงานนี้มีวัตถุประสงค์เพื่อพัฒนาเว็บแอปพลิเคชันสำหรับจัดการธุรกิจค้าส่ง 123พาณิชย์ปลีกส่ง ให้ผู้เกี่ยวข้องทำงานบนข้อมูลสินค้า สินค้าคงเหลือ คำสั่งซื้อ ใบเสร็จรายวัน และข้อมูลการนำส่งชุดเดียวกัน ระบบกำหนดผู้ใช้งานเป็น 4 บทบาท ได้แก่ เจ้าของร้าน ผู้ดูแลคลังสินค้า พนักงาน และผู้ค้าปลีก เพื่อให้ขอบเขตการเข้าถึงข้อมูลเหมาะสมกับหน้าที่ของแต่ละฝ่าย",
            "ผู้ค้าปลีกสามารถสมัครเข้าใช้ เลือกสินค้าตามประเภท เลือกจำนวนแบบกล่องหรือแบบแพ็ค ตรวจสอบตะกร้าสินค้า และยืนยันคำสั่งซื้อได้ ระบบตรวจสอบราคาและสินค้าคงเหลือก่อนบันทึกคำสั่งซื้อ สร้างหรือรวมข้อมูลใบเสร็จรายวันของผู้ค้าปลีก และปรับจำนวนคงเหลือด้วยกระบวนการธุรกรรมเดียวกัน ฝั่งพนักงานจัดการสถานะคำสั่งซื้อตามลำดับ PENDING, PACKED, DELIVERING และ COMPLETED พร้อมแชร์ตำแหน่งระหว่างนำส่งเมื่อผู้ใช้อนุญาต ส่วนเจ้าของร้านและผู้ดูแลคลังสินค้าจัดการสินค้า ประเภทสินค้า และการปรับสต็อกได้ตามสิทธิ์",
            "ระบบพัฒนาด้วย Next.js, React, TypeScript และ Tailwind CSS เชื่อมต่อฐานข้อมูล Supabase PostgreSQL ผ่านส่วนประมวลผลฝั่งเซิร์ฟเวอร์ ใช้การยืนยันตัวตนแบบกำหนดเองด้วยรหัสผ่านที่ผ่านการแฮชและ session cookie ที่ลงลายมือชื่อ รายงานฉบับนี้อธิบายการวิเคราะห์ ออกแบบ พัฒนา และทดสอบการทำงานตามขอบเขตของระบบ โดยมุ่งให้เป็นเอกสารประกอบโครงงานที่สอดคล้องกับความสามารถของเว็บแอปพลิเคชันจริง",
        ]
        keyword = "คำสำคัญ: เว็บแอปพลิเคชัน, ระบบค้าส่ง, การจัดการสต็อก, การจัดส่งสินค้า, การกำหนดสิทธิ์ผู้ใช้"
    else:
        paragraphs = [
            "This project develops a web application for 123 Phanit Retail and Wholesale Management. The system brings product, inventory, order, daily-receipt, and delivery information into one shared data source. Four roles are provided: owner, inventory administrator, employee, and retailer, allowing each participant to work within an appropriate access boundary.",
            "Retailers can register, browse products by category, select box or pack quantities, review a cart, and confirm an order. Before an order is saved, the system validates prices and available stock, creates or reuses the retailer's daily receipt, and updates stock within one transaction. Employees progress assigned work through PENDING, PACKED, DELIVERING, and COMPLETED, and can share their location while delivering. Owners and administrators manage catalog data and stock according to their permissions.",
            "The application is implemented with Next.js, React, TypeScript, Tailwind CSS, and Supabase PostgreSQL. Custom authentication uses hashed passwords and signed session cookies. This report presents the analysis, design, implementation, and functional test approach based on the actual scope of the web application.",
        ]
        keyword = "Keywords: web application, wholesale management, inventory control, delivery tracking, role-based access"
    for t in paragraphs:
        add_text(doc, t, size=16, align=WD_ALIGN_PARAGRAPH.JUSTIFY if english else WD_ALIGN_PARAGRAPH.LEFT, space_after=7)
    add_text(doc, keyword, size=14, bold=True, first_indent=False, space_before=8)


def front_matter(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(14)
    set_run_font(p.add_run("สารบัญ"), 22, bold=True, color="000000")
    items = [
        "บทคัดย่อ", "Abstract", "สารบัญ", "สารบัญตาราง", "สารบัญภาพ",
        "บทที่ 1 บทนำ", "บทที่ 2 เอกสารและงานที่เกี่ยวข้อง", "บทที่ 3 วิธีดำเนินการวิจัย",
        "บทที่ 4 ผลการดำเนินงาน", "บทที่ 5 สรุปผลและข้อเสนอแนะ", "บรรณานุกรม", "ภาคผนวก",
    ]
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Mm(10)
        p.paragraph_format.space_after = Pt(4)
        set_run_font(p.add_run(item), 16)
    doc.add_page_break()
    for heading, labels in [
        ("สารบัญตาราง", [
            "ตารางที่ 1.1 ขอบเขตการใช้งานตามบทบาท", "ตารางที่ 1.2 เครื่องมือและเทคโนโลยีหลัก", "ตารางที่ 1.3 ระยะเวลาการดำเนินงาน",
            "ตารางที่ 2.1 เทคโนโลยีและบทบาทในการพัฒนา", "ตารางที่ 3.1 พจนานุกรมข้อมูลหลัก",
            "ตารางที่ 3.2 เงื่อนไขสิทธิ์ตามบทบาท", "ตารางที่ 4.1 กรณีทดสอบเชิงหน้าที่",
        ]),
        ("สารบัญภาพ", [
            "ภาพที่ 3.1 ภาพรวมสถาปัตยกรรมระบบ", "ภาพที่ 3.2 การไหลของคำสั่งซื้อและการจัดส่ง",
            "ภาพที่ 3.3 แบบจำลองข้อมูลหลัก", "ภาพที่ 3.4 ขอบเขตการใช้งานตามบทบาท",
        ]),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(14)
        set_run_font(p.add_run(heading), 22, bold=True, color="000000")
        for label in labels:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Mm(10)
            p.paragraph_format.space_after = Pt(4)
            set_run_font(p.add_run(label), 16)
        doc.add_page_break()


def chapter1(doc):
    add_heading(doc, "บทนำ", level=1, number="บทที่ 1")
    add_heading(doc, "ที่มาและความสำคัญ", level=2, number="1.1")
    add_text(doc, "ธุรกิจค้าส่งต้องรับข้อมูลการสั่งซื้อ ตรวจสอบสินค้า เตรียมการจัดส่ง และตอบคำถามเกี่ยวกับสถานะงานอย่างต่อเนื่อง เมื่อข้อมูลดังกล่าวกระจายอยู่ในหลายช่องทาง การตรวจสอบจำนวนสินค้าและความคืบหน้าของออเดอร์อาจล่าช้าหรือเกิดความซ้ำซ้อนได้ ร้าน 123พาณิชย์ปลีกส่งจึงมีความต้องการระบบที่รวมข้อมูลการทำงานสำคัญไว้ในพื้นที่เดียว เพื่อให้ผู้เกี่ยวข้องเข้าถึงข้อมูลได้ตรงตามหน้าที่")
    add_text(doc, "โครงงานนี้พัฒนาเว็บแอปพลิเคชันที่เชื่อมเจ้าของร้าน ผู้ดูแลคลังสินค้า พนักงาน และผู้ค้าปลีกเข้ากับฐานข้อมูลกลาง ระบบรองรับการจัดการสินค้าและหมวดหมู่ การปรับสต็อก การเลือกสินค้าเป็นกล่องหรือแพ็ค การยืนยันคำสั่งซื้อ การสร้างใบเสร็จรายวัน และการติดตามสถานะการนำส่ง โดยพนักงานสามารถแชร์ตำแหน่งผ่านอุปกรณ์ของตนได้เมื่ออยู่ระหว่างนำส่ง")
    add_heading(doc, "วัตถุประสงค์", level=2, number="1.2")
    objectives = [
        "เพื่อพัฒนาเว็บแอปพลิเคชันสำหรับจัดการข้อมูลสินค้า สต็อก คำสั่งซื้อ ใบเสร็จรายวัน และการนำส่งของร้านค้าส่ง",
        "เพื่อกำหนดสิทธิ์การใช้งานที่เหมาะสมสำหรับเจ้าของร้าน ผู้ดูแลคลังสินค้า พนักงาน และผู้ค้าปลีก",
        "เพื่อให้ผู้ค้าปลีกสั่งสินค้าแบบกล่องหรือแบบแพ็ค และตรวจสอบข้อมูลการซื้อของตนเองได้",
        "เพื่อสนับสนุนการควบคุมสต็อก การป้องกันการบันทึกคำสั่งซื้อซ้ำ และการติดตามงานนำส่ง",
    ]
    for i, item in enumerate(objectives, 1):
        add_text(doc, f"{i}. {item}", first_indent=False, space_after=3)
    add_heading(doc, "ขอบเขตของโครงงาน", level=2, number="1.3")
    add_heading(doc, "ความสามารถของระบบ / เนื้อหาที่ใช้", level=2, number="1.3.1")
    add_scope_group(doc, 1, "ฟังก์ชันพื้นฐานของระบบ (System Core Functions)", [
        "ระบบสมัครสมาชิกและเข้าสู่ระบบด้วยหมายเลขโทรศัพท์และรหัสผ่าน โดยรหัสผ่านถูกแฮชก่อนจัดเก็บ และ session ใช้ cookie แบบ httpOnly ที่มีการลงลายมือชื่อ เพื่อช่วยลดการเปิดเผยข้อมูลการยืนยันตัวตน",
        "แบ่งผู้ใช้งานออกเป็น 4 บทบาท ได้แก่ OWNER (เจ้าของร้าน), ADMIN (ผู้ดูแลคลัง), EMPLOYEE (พนักงาน) และ RETAILER (ผู้ค้าปลีก) พร้อมตรวจสอบสิทธิ์ทั้งในหน้าจอและส่วนให้บริการข้อมูล",
        "ระบบแสดงสินค้าเป็นหมวดหมู่ พร้อมค้นหาและกรองรายการ แสดงราคาต่อกล่อง ราคาต่อแพ็ค และจำนวนคงเหลือ เพื่อให้ผู้ค้าปลีกตัดสินใจเลือกสินค้าได้จากข้อมูลปัจจุบัน",
        "จัดเก็บข้อมูลธุรกิจในฐานข้อมูลกลาง โดยเก็บชื่อสินค้า ราคา และยอดรายการไว้กับข้อมูลธุรกรรม เพื่อให้ตรวจสอบประวัติได้แม้สินค้าถูกนำออกจากรายการสินค้าปัจจุบัน",
    ])
    add_scope_group(doc, 2, "ฟังก์ชันสำหรับเจ้าของร้าน (OWNER)", [
        "เพิ่ม แก้ไข และลบข้อมูลสินค้า รวมทั้งกำหนดราคาแบบกล่องหรือแบบแพ็ค และตรวจสอบจำนวนสินค้าคงเหลือ",
        "จัดการหมวดหมู่สินค้า ได้แก่ เพิ่ม เปลี่ยนชื่อ เปิดหรือปิดการใช้งาน และลบได้ตามข้อกำหนดความสัมพันธ์ของข้อมูล",
        "ตรวจสอบภาพรวมคำสั่งซื้อ ใบเสร็จรายวัน ข้อมูลยอดขาย และสถานะการนำส่ง เพื่อใช้ประกอบการติดตามการดำเนินงานของร้าน",
        "เปิดดูตำแหน่งล่าสุดของงานนำส่งผ่านแผนที่ OpenStreetMap และเปิดพิกัดต่อใน Google Maps ได้",
    ])
    add_scope_group(doc, 3, "ฟังก์ชันสำหรับผู้ดูแลคลัง (ADMIN)", [
        "ตรวจสอบรายการสินค้าและปรับเพิ่มหรือลดจำนวนคงเหลือตามสิทธิ์ที่ได้รับ เพื่อบันทึกการเปลี่ยนแปลงของสินค้าคงคลัง",
        "ขอบเขตของบทบาทนี้มุ่งที่งานคลังสินค้า จึงไม่รวมการจัดการบัญชีผู้ใช้ การจัดการสินค้าและหมวดหมู่ หรือการจัดการคำสั่งซื้อแทนเจ้าของร้าน",
    ])
    add_scope_group(doc, 4, "ฟังก์ชันสำหรับพนักงาน (EMPLOYEE)", [
        "ตรวจสอบคำสั่งซื้อและดำเนินงานตามลำดับสถานะ PENDING → PACKED → DELIVERING → COMPLETED โดยระบบตรวจสอบลำดับก่อนบันทึกผล",
        "เมื่อพนักงานเริ่มเปลี่ยนสถานะของคำสั่งซื้อที่ยังไม่มีผู้รับผิดชอบ ระบบจะกำหนดผู้ดำเนินงานรายนั้นและป้องกันไม่ให้พนักงานคนอื่นรับงานเดียวกันซ้ำ",
        "เมื่ออยู่ในสถานะ DELIVERING พนักงานสามารถอนุญาตให้เบราว์เซอร์ส่งพิกัดตำแหน่งล่าสุด ระบบจะส่งข้อมูลเมื่อเคลื่อนที่อย่างน้อย 25 เมตรหรือเมื่อครบช่วงเวลา 15 วินาที โดยเว้นระยะการส่งอย่างน้อย 5 วินาที",
    ])
    add_scope_group(doc, 5, "ฟังก์ชันสำหรับผู้ค้าปลีก (RETAILER)", [
        "สมัครบัญชีและเข้าสู่ระบบเพื่อใช้งานข้อมูลของร้านตนเอง ค้นหาและกรองสินค้า เลือกจำนวนแบบกล่องหรือแบบแพ็ค แล้วเพิ่มรายการลงตะกร้าสินค้า",
        "ยืนยันคำสั่งซื้อโดยระบบตรวจสอบราคาและจำนวนคงเหลือ ล็อกข้อมูลสต็อก สร้างคำสั่งซื้อและรายการสินค้า สร้างหรือรวมใบเสร็จรายวันตามวันในเขตเวลา Asia/Bangkok และหักสต็อกในธุรกรรมเดียวกัน",
        "ใช้ request id ป้องกันการยืนยันคำสั่งซื้อซ้ำ ดูประวัติใบเสร็จรายวัน และตรวจสอบสถานะหรือตำแหน่งล่าสุดของคำสั่งซื้อของร้านตนเองได้",
    ])
    add_scope_group(doc, 6, "ข้อมูล เทคโนโลยี และข้อจำกัดของระบบ", [
        "ข้อมูลหลักประกอบด้วยบัญชีผู้ใช้และข้อมูลร้านค้า หมวดหมู่สินค้า สินค้า คำสั่งซื้อและรายการสินค้า ใบเสร็จรายวันและรายการใบเสร็จ รวมถึงตำแหน่งการนำส่งล่าสุด",
        "พัฒนาด้วย Next.js 16, React 19, TypeScript และ Tailwind CSS 4 เชื่อมต่อ Supabase PostgreSQL ผ่าน Next.js Route Handlers และใช้ Browser Geolocation ร่วมกับ OpenStreetMap และลิงก์ Google Maps สำหรับข้อมูลตำแหน่ง",
        "การแสดงข้อมูลหน้าจอใช้การดึงข้อมูลเป็นระยะ จึงเป็นการติดตามแบบใกล้เคียงเรียลไทม์ ไม่ใช่ระบบ realtime เต็มรูปแบบ",
        "ไม่ครอบคลุมระบบรับชำระเงินออนไลน์ ระบบภาษีหรือบัญชี การแปลงหน่วยระหว่างกล่องกับแพ็ค การวางแผนเส้นทางอัตโนมัติ ประวัติเส้นทางย้อนหลัง การแจ้งเตือนแบบ push notification และหน้าจอจัดการบัญชีพนักงาน",
    ])
    doc.add_page_break()
    add_table(doc, "ตารางที่ 1.1 ขอบเขตการใช้งานตามบทบาท", ["บทบาท", "ขอบเขตการทำงาน"], [
        ["OWNER\nเจ้าของร้าน", "จัดการสินค้า หมวดหมู่ และสต็อก ดูภาพรวมใบเสร็จรายวัน และติดตามสถานะการนำส่ง"],
        ["ADMIN\nผู้ดูแลคลัง", "ตรวจสอบข้อมูลสินค้าและปรับเพิ่มหรือลดจำนวนคงเหลือตามสิทธิ์"],
        ["EMPLOYEE\nพนักงาน", "รับผิดชอบคำสั่งซื้อ เปลี่ยนสถานะตามลำดับ และแชร์พิกัดเมื่ออยู่ในสถานะ DELIVERING"],
        ["RETAILER\nผู้ค้าปลีก", "สมัครบัญชี ค้นหาและสั่งสินค้าแบบกล่องหรือแพ็ค ดูใบเสร็จรายวัน และติดตามคำสั่งซื้อของร้านตนเอง"],
    ], [Mm(42), Mm(118)], 10.3)
    add_heading(doc, "เครื่องมือที่ใช้ในการพัฒนา", level=2, number="1.4")
    add_table(doc, "ตารางที่ 1.2 เครื่องมือและเทคโนโลยีหลัก", ["กลุ่ม", "รายการ", "การใช้งาน"], [
        ["ส่วนติดต่อผู้ใช้", "Next.js 16, React 19, TypeScript, Tailwind CSS 4", "สร้างหน้าเว็บ แดชบอร์ด และองค์ประกอบโต้ตอบ"],
        ["ส่วนประมวลผล", "Next.js Route Handlers", "ตรวจสอบ session บทบาท และประมวลผลคำขอจากหน้าเว็บ"],
        ["ฐานข้อมูล", "Supabase PostgreSQL, REST/RPC", "จัดเก็บข้อมูลธุรกิจและดำเนินธุรกรรมคำสั่งซื้อ"],
        ["แผนที่", "Browser Geolocation, OpenStreetMap, Google Maps link", "รับพิกัดและแสดงตำแหน่งล่าสุดระหว่างนำส่ง"],
    ], [Mm(34), Mm(55), Mm(71)], 9.5)
    add_heading(doc, "แผนการดำเนินงาน", level=2, number="1.5", page_break_before=True)
    add_schedule_table(doc)
    add_heading(doc, "ประโยชน์ที่คาดว่าจะได้รับ", level=2, number="1.6")
    for i, item in enumerate([
        "ช่วยให้ข้อมูลสินค้า สต็อก คำสั่งซื้อ และการนำส่งอยู่ในแหล่งข้อมูลเดียวกัน",
        "ลดขั้นตอนการสอบถามข้อมูลซ้ำ และช่วยให้ผู้ค้าปลีกตรวจสอบการซื้อของตนเองได้",
        "ลดความเสี่ยงสต็อกติดลบและการสร้างคำสั่งซื้อซ้ำจากการยืนยันซ้ำ",
        "ช่วยให้เจ้าของร้านติดตามภาพรวมการขายและความคืบหน้าการนำส่งได้สะดวกขึ้น",
    ], 1):
        add_text(doc, f"{i}. {item}", first_indent=False, space_after=3)


def chapter2(doc):
    add_heading(doc, "เอกสารและงานที่เกี่ยวข้อง", level=1, number="บทที่ 2", page_break_before=True)
    add_heading(doc, "แนวคิดการพัฒนาเว็บแอปพลิเคชัน", level=2, number="2.1")
    add_text(doc, "เว็บแอปพลิเคชันเป็นรูปแบบระบบสารสนเทศที่ผู้ใช้เข้าถึงผ่านเว็บเบราว์เซอร์ ช่วยให้ผู้ใช้หลายบทบาทใช้งานข้อมูลชุดเดียวกันได้โดยไม่ต้องติดตั้งโปรแกรมเฉพาะเครื่อง การออกแบบระบบค้าส่งจึงต้องคำนึงถึงความถูกต้องของข้อมูล ความชัดเจนของสิทธิ์ และความต่อเนื่องของกระบวนการตั้งแต่การเลือกสินค้าไปจนถึงการส่งมอบ")
    add_heading(doc, "การควบคุมสิทธิ์และการยืนยันตัวตน", level=2, number="2.2")
    add_text(doc, "ระบบกำหนดสิทธิ์ตามบทบาททั้งในระดับหน้าแดชบอร์ดและจุดให้บริการ API เพื่อให้ผู้ใช้เห็นและดำเนินการเฉพาะข้อมูลที่เกี่ยวข้องกับหน้าที่ของตน การเข้าสู่ระบบใช้รหัสผ่านที่ผ่านการแฮชร่วมกับ salt และสร้าง session ที่มีการลงลายมือชื่อ HMAC ใน cookie แบบ httpOnly แนวทางดังกล่าวช่วยลดการเปิดเผยรหัสผ่านจริงและลดการเข้าถึงข้อมูลโดยไม่ได้รับอนุญาต")
    add_heading(doc, "การจัดการธุรกรรมคำสั่งซื้อ", level=2, number="2.3")
    add_text(doc, "การยืนยันคำสั่งซื้อเป็นจุดที่เกี่ยวข้องกับข้อมูลสำคัญหลายส่วน ได้แก่ ราคา สต็อก คำสั่งซื้อ รายการสินค้า และใบเสร็จรายวัน ระบบจึงดำเนินการตรวจสอบข้อมูลและปรับสต็อกภายในกระบวนการธุรกรรมเดียวกัน พร้อมใช้กลไกล็อกข้อมูลและ request id เพื่อป้องกันกรณีผู้ใช้กดยืนยันซ้ำ ผลลัพธ์คือระบบไม่สร้างคำสั่งซื้อหรือหักสต็อกซ้ำจากคำขอเดิม")
    add_heading(doc, "เทคโนโลยีที่เกี่ยวข้อง", level=2, number="2.4")
    add_table(doc, "ตารางที่ 2.1 เทคโนโลยีและบทบาทในการพัฒนา", ["เทคโนโลยี", "บทบาทในระบบ"], [
        ["Next.js", "จัดการโครงสร้างเว็บแอปพลิเคชัน หน้าจอ และ route handlers ฝั่งเซิร์ฟเวอร์"],
        ["React", "พัฒนาองค์ประกอบส่วนติดต่อผู้ใช้และสถานะของหน้าแดชบอร์ด"],
        ["TypeScript", "ระบุชนิดข้อมูล ลดความคลาดเคลื่อนในการเชื่อมข้อมูลระหว่างส่วนต่าง ๆ"],
        ["Supabase PostgreSQL", "จัดเก็บข้อมูลธุรกิจ รองรับ REST/RPC และข้อกำหนดความถูกต้องของฐานข้อมูล"],
        ["Tailwind CSS", "กำหนดรูปแบบหน้าจอให้ตอบสนองต่อขนาดอุปกรณ์และใช้สไตล์ร่วมกัน"],
    ], [Mm(48), Mm(112)], 10.0)
    add_heading(doc, "เอกสารอ้างอิงและมาตรฐานที่ใช้ประกอบ", level=2, number="2.5")
    add_text(doc, "เอกสารประกอบการพัฒนาระบบอ้างอิงแนวทางจากเอกสารทางการของ Next.js, React และ Supabase เพื่อให้โครงสร้างการพัฒนาและการเชื่อมต่อฐานข้อมูลสอดคล้องกับการใช้งานของเทคโนโลยี นอกจากนี้ได้ใช้แนวคิดเรื่องการควบคุมสิทธิ์ การตรวจสอบข้อมูล และความถูกต้องของธุรกรรมเป็นแนวทางในการออกแบบส่วนที่เกี่ยวข้องกับสต็อกและคำสั่งซื้อ")
    add_text(doc, "สำหรับข้อมูลธุรกิจ ระบบใช้ฐานข้อมูลเชิงสัมพันธ์เพื่อกำหนดความเชื่อมโยงระหว่างผู้ใช้ สินค้า รายการคำสั่งซื้อ และใบเสร็จ โดยใช้ข้อกำหนด เช่น foreign key และ check constraint เป็นส่วนช่วยตรวจสอบความสอดคล้องของข้อมูล ส่วนการจำกัดการเข้าถึงข้อมูลใช้บทบาทในระบบร่วมกับนโยบายระดับฐานข้อมูล")
    add_text(doc, "การออกแบบดังกล่าวเน้นให้ข้อมูลจากรายการธุรกรรมสามารถตรวจสอบย้อนหลังได้ โดยเก็บชื่อสินค้า ราคาต่อหน่วย และยอดของรายการไว้กับข้อมูลออเดอร์หรือใบเสร็จในเวลาที่เกิดรายการ จึงยังอ้างอิงประวัติได้แม้สินค้าเดิมถูกนำออกจากรายการสินค้าปัจจุบัน")


def chapter3(doc):
    add_heading(doc, "วิธีดำเนินการวิจัย", level=1, number="บทที่ 3", page_break_before=True)
    add_heading(doc, "การวิเคราะห์ระบบ", level=2, number="3.1")
    add_text(doc, "เริ่มจากการวิเคราะห์ผู้มีส่วนเกี่ยวข้องและข้อมูลที่ต้องใช้ร่วมกัน จากนั้นกำหนดบทบาท แยกความสามารถของแต่ละบทบาท และออกแบบความสัมพันธ์ของข้อมูล ระบบให้ผู้ค้าปลีกทำรายการด้วยตนเอง ขณะที่พนักงานจัดการงานที่ได้รับมอบหมาย และเจ้าของร้านตรวจสอบข้อมูลภาพรวม การออกแบบดังกล่าวช่วยลดการส่งต่อข้อมูลด้วยวิธีที่ไม่เป็นระบบ")
    add_figure(doc, ASSETS / "architecture.png", "ภาพที่ 3.1 ภาพรวมสถาปัตยกรรมระบบ")
    add_text(doc, "ผู้ใช้เข้าถึงส่วนติดต่อบนเว็บผ่านเว็บเบราว์เซอร์ คำขอที่เกี่ยวข้องกับข้อมูลสำคัญส่งผ่าน Next.js Route Handlers ซึ่งตรวจสอบ session และบทบาทก่อนเชื่อมต่อ Supabase PostgreSQL ฝั่งเซิร์ฟเวอร์จึงเป็นจุดควบคุมการเข้าถึงข้อมูลและการใช้ service key ส่วนแผนที่ใช้ OpenStreetMap สำหรับแสดงผล และมีลิงก์สำหรับเปิดตำแหน่งต่อใน Google Maps")
    add_heading(doc, "การออกแบบกระบวนการคำสั่งซื้อ", level=2, number="3.2")
    add_figure(doc, ASSETS / "order-flow.png", "ภาพที่ 3.2 การไหลของคำสั่งซื้อและการจัดส่ง")
    add_text(doc, "ผู้ค้าปลีกเลือกสินค้าพร้อมจำนวนแบบกล่องหรือแบบแพ็คแล้วส่งคำขอยืนยัน ระบบตรวจสอบสินค้า ราคา และจำนวนคงเหลือก่อนสร้างข้อมูลคำสั่งซื้อและรายการสินค้า หลังจากนั้นจะสร้างหรือรวมข้อมูลใบเสร็จรายวันของผู้ค้าปลีกตามวันในเขตเวลา Asia/Bangkok และหักสต็อกในกระบวนการเดียวกัน พนักงานดำเนินงานตามสถานะที่กำหนดเพื่อทำให้การส่งต่อความรับผิดชอบตรวจสอบได้")
    add_heading(doc, "การออกแบบข้อมูล", level=2, number="3.3")
    add_figure(doc, ASSETS / "data-model.png", "ภาพที่ 3.3 แบบจำลองข้อมูลหลักของระบบ")
    add_table(doc, "ตารางที่ 3.1 พจนานุกรมข้อมูลหลัก", ["ตารางข้อมูล", "หน้าที่"], [
        ["profiles", "จัดเก็บบัญชี ชื่อร้าน ข้อมูลติดต่อ และบทบาทของผู้ใช้"],
        ["product_categories", "จัดเก็บหมวดหมู่สินค้าและสถานะการใช้งาน"],
        ["products", "จัดเก็บสินค้า ราคาแบบกล่อง/แพ็ค และจำนวนคงเหลือ"],
        ["orders และ order_items", "จัดเก็บคำสั่งซื้อ สถานะ ผู้รับผิดชอบ และรายการสินค้าของแต่ละออเดอร์"],
        ["daily_receipts และ receipt_items", "จัดเก็บใบเสร็จรวมรายวันและรายการสินค้า เพื่อให้ร้านหนึ่งมีใบเสร็จหนึ่งรายการต่อวัน"],
        ["delivery_tracking", "จัดเก็บพิกัดล่าสุดของพนักงานสำหรับคำสั่งซื้อที่เกี่ยวข้อง"],
    ], [Mm(62), Mm(98)], 9.7)
    add_heading(doc, "การออกแบบสิทธิ์การใช้งาน", level=2, number="3.4")
    add_figure(doc, ASSETS / "roles.png", "ภาพที่ 3.4 ขอบเขตการใช้งานตามบทบาท")
    add_table(doc, "ตารางที่ 3.2 เงื่อนไขการเข้าถึงตามบทบาท", ["การดำเนินการ", "บทบาทที่อนุญาต"], [
        ["เพิ่ม/ลบสินค้า และจัดการประเภทสินค้า", "OWNER"],
        ["ปรับจำนวนสต็อก", "OWNER, ADMIN"],
        ["ยืนยันคำสั่งซื้อ", "RETAILER"],
        ["เปลี่ยนสถานะออเดอร์และส่งตำแหน่ง", "EMPLOYEE"],
        ["ดูคำสั่งซื้อและใบเสร็จของร้านตนเอง", "RETAILER"],
        ["ดูภาพรวมรายรับและการนำส่ง", "OWNER"],
    ], [Mm(97), Mm(63)], 10.0)
    add_heading(doc, "การพัฒนาและการทดสอบ", level=2, number="3.5")
    add_text(doc, "การพัฒนาทำเป็นส่วนย่อย ได้แก่ ส่วนยืนยันตัวตน แดชบอร์ดตามบทบาท การจัดการสินค้า กระบวนการยืนยันคำสั่งซื้อ การจัดการสถานะ และการติดตามพิกัด จากนั้นทดสอบกรณีที่เกี่ยวข้องกับสิทธิ์ การเปลี่ยนสถานะ การตรวจสอบสต็อก และการป้องกันการส่งคำขอซ้ำ การทดสอบใช้ข้อมูลตัวอย่างและตรวจสอบผลลัพธ์จากฐานข้อมูลร่วมกับพฤติกรรมของหน้าเว็บ")


def chapter4(doc):
    add_heading(doc, "ผลการดำเนินงาน", level=1, number="บทที่ 4", page_break_before=True)
    add_heading(doc, "ผลการพัฒนาระบบ", level=2, number="4.1")
    add_text(doc, "ระบบที่พัฒนาสามารถแยกหน้าการทำงานตามบทบาทได้ เจ้าของร้านเข้าถึงข้อมูลภาพรวมสินค้า สต็อก ใบเสร็จ และสถานะการจัดส่ง ผู้ดูแลคลังปรับสต็อกตามสิทธิ์ พนักงานจัดการลำดับสถานะคำสั่งซื้อและพิกัดระหว่างนำส่ง ส่วนผู้ค้าปลีกเลือกสินค้าและยืนยันคำสั่งซื้อได้ด้วยตนเอง ระบบจัดเก็บข้อมูลธุรกรรมเป็นประวัติ เพื่อให้ตรวจสอบย้อนหลังได้")
    add_heading(doc, "การทำงานของใบเสร็จรายวันและสต็อก", level=2, number="4.2")
    add_text(doc, "เมื่อผู้ค้าปลีกยืนยันคำสั่งซื้อ ระบบจะตรวจสอบสินค้าคงเหลือและข้อมูลราคาก่อนสร้าง order และ order items พร้อมกับสร้างหรือใช้ daily receipt เดิมของผู้ค้าปลีกรายนั้นในวันเดียวกัน รายการสินค้าจะถูกเพิ่มลง receipt items และสต็อกลดลงทันทีในธุรกรรมเดียวกัน กลไกดังกล่าวทำให้จำนวนคงเหลือสะท้อนสินค้าที่ถูกยืนยันคำสั่งซื้อแล้ว และช่วยลดความเสี่ยงที่รายการหลายรายการจะแย่งข้อมูลสต็อกชุดเดียวกัน")
    add_heading(doc, "ผลการทดสอบเชิงหน้าที่", level=2, number="4.3")
    add_table(doc, "ตารางที่ 4.1 กรณีทดสอบเชิงหน้าที่ของระบบ", ["กรณีทดสอบ", "เงื่อนไข", "ผลที่คาดหวัง"], [
        ["เข้าสู่ระบบ", "ใช้ข้อมูลผู้ใช้ที่ถูกต้องและบทบาทต่างกัน", "เข้าสู่แดชบอร์ดที่ตรงกับบทบาท และไม่เข้าถึงหน้าที่ไม่ได้รับสิทธิ์"],
        ["ยืนยันคำสั่งซื้อ", "RETAILER เลือกสินค้าที่มีสต็อกเพียงพอ", "สร้าง order/order_items สร้างหรือรวมใบเสร็จรายวัน และลดสต็อก"],
        ["ยืนยันคำขอซ้ำ", "ส่ง request id เดิมซ้ำ", "ได้ผลลัพธ์ของคำขอเดิมโดยไม่สร้างออเดอร์และไม่หักสต็อกซ้ำ"],
        ["สต็อกไม่เพียงพอ", "เลือกจำนวนมากกว่าคงเหลือ", "ระบบปฏิเสธรายการและคงข้อมูลสต็อกเดิม"],
        ["เปลี่ยนสถานะ", "EMPLOYEE เปลี่ยนสถานะตามลำดับ", "อนุญาตเฉพาะลำดับ PENDING → PACKED → DELIVERING → COMPLETED"],
        ["ตำแหน่งนำส่ง", "EMPLOYEE อยู่สถานะ DELIVERING และอนุญาต GPS", "บันทึกตำแหน่งล่าสุดเพื่อให้เจ้าของและผู้ค้าปลีกที่เกี่ยวข้องตรวจสอบได้"],
    ], [Mm(40), Mm(56), Mm(64)], 9.0)
    add_heading(doc, "การวิเคราะห์ผล", level=2, number="4.4")
    add_text(doc, "ผลการออกแบบทำให้ข้อมูลหลักของธุรกิจถูกจัดเก็บและใช้งานร่วมกันในระบบเดียว ผู้ค้าปลีกสามารถตรวจสอบตะกร้า ใบเสร็จ และออเดอร์ของร้านตนเอง ขณะที่ผู้ปฏิบัติงานเห็นข้อมูลที่เกี่ยวข้องกับหน้าที่โดยตรง การใช้ transaction, row lock และ idempotency key ที่จุดยืนยันคำสั่งซื้อช่วยให้กระบวนการสำคัญมีเงื่อนไขตรวจสอบชัดเจน ทั้งนี้การแสดงข้อมูลเป็นแบบใกล้เคียงเรียลไทม์จากการดึงข้อมูลเป็นระยะ จึงไม่ควรตีความว่าเป็นระบบ realtime เต็มรูปแบบ")
    add_heading(doc, "ข้อจำกัดของระบบ", level=2, number="4.5")
    add_text(doc, "ระบบปัจจุบันมุ่งเน้นการจัดการออเดอร์และการประสานงานภายในร้าน จึงยังไม่รวมการรับชำระเงินออนไลน์ การวางแผนเส้นทางจัดส่งอัตโนมัติ และการแจ้งเตือนแบบ push notification การแชร์ตำแหน่งขึ้นอยู่กับการอนุญาตตำแหน่งและการเชื่อมต่ออินเทอร์เน็ตของอุปกรณ์พนักงาน")


def chapter5(doc):
    add_heading(doc, "สรุปผลและข้อเสนอแนะ", level=1, number="บทที่ 5", page_break_before=True)
    add_heading(doc, "สรุปผล", level=2, number="5.1")
    add_text(doc, "โครงงานนี้ได้พัฒนาเว็บแอปพลิเคชันระบบจัดการค้าส่งสำหรับ 123พาณิชย์ปลีกส่ง โดยครอบคลุมการจัดการสินค้า สต็อก คำสั่งซื้อ ใบเสร็จรายวัน และการนำส่ง ระบบแบ่งสิทธิ์สำหรับเจ้าของร้าน ผู้ดูแลคลังสินค้า พนักงาน และผู้ค้าปลีก เพื่อให้การเข้าถึงข้อมูลสอดคล้องกับความรับผิดชอบของแต่ละฝ่าย")
    add_text(doc, "จุดสำคัญของระบบคือกระบวนการยืนยันคำสั่งซื้อที่ตรวจสอบสต็อกและข้อมูลราคา สร้างข้อมูลธุรกรรมที่เกี่ยวข้อง และปรับสต็อกในขั้นตอนเดียวกันพร้อมการป้องกันคำขอซ้ำ ขณะเดียวกันระบบสนับสนุนการติดตามสถานะของพนักงานตามลำดับงาน และให้ผู้เกี่ยวข้องตรวจสอบตำแหน่งล่าสุดได้เมื่ออยู่ในช่วงนำส่ง")
    add_heading(doc, "อภิปรายผล", level=2, number="5.2")
    add_text(doc, "การรวมข้อมูลหลักไว้ในระบบเดียวช่วยให้บทบาทต่าง ๆ อ้างอิงข้อมูลชุดเดียวกัน แทนการส่งข้อมูลผ่านหลายช่องทาง การกำหนดขอบเขตข้อมูลตามบทบาทช่วยลดการเข้าถึงเกินความจำเป็น ส่วนการเก็บข้อมูลรายการสินค้าและราคาขณะทำธุรกรรมช่วยรักษาประวัติ แม้สินค้าบางรายการจะถูกนำออกจากรายการสินค้าปัจจุบัน")
    add_heading(doc, "ข้อเสนอแนะ", level=2, number="5.3")
    for i, item in enumerate([
        "เพิ่มการแจ้งเตือนเมื่อมีคำสั่งซื้อใหม่ เปลี่ยนสถานะ หรือสินค้าใกล้หมด โดยพิจารณาช่องทางที่เหมาะกับผู้ใช้",
        "ต่อยอดการแสดงรายงานเชิงวิเคราะห์ เช่น สินค้าขายดี ยอดขายตามช่วงเวลา และแนวโน้มสต็อก",
        "ศึกษาการวางแผนเส้นทางและการจัดกลุ่มออเดอร์สำหรับการนำส่งในอนาคต",
        "เก็บข้อมูลการใช้งานและความพึงพอใจจากผู้ใช้จริง เพื่อประเมินความเหมาะสมของส่วนติดต่อผู้ใช้ในรอบถัดไป",
    ], 1):
        add_text(doc, f"{i}. {item}", first_indent=False, space_after=3)


def references_and_appendices(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.page_break_before = True
    p.paragraph_format.space_after = Pt(15)
    set_run_font(p.add_run("บรรณานุกรม"), 22, bold=True, color="000000")
    entries = [
        "Next.js. (2026). Next.js Documentation. https://nextjs.org/docs",
        "React. (2026). React Documentation. https://react.dev",
        "Supabase. (2026). Supabase Documentation. https://supabase.com/docs",
        "OWASP Foundation. (2026). OWASP Application Security Verification Standard. https://owasp.org/www-project-application-security-verification-standard/",
        "ซอร์สโค้ดระบบ 123พาณิชย์ปลีกส่ง. (2569). เอกสารและโครงสร้างโปรแกรมภายในโครงการ.",
    ]
    for e in entries:
        add_text(doc, e, size=14, first_indent=False, space_after=6)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.page_break_before = True
    p.paragraph_format.space_after = Pt(14)
    set_run_font(p.add_run("ภาคผนวก ก\nคู่มือย่อการใช้งานตามบทบาท"), 22, bold=True, color="000000")
    instructions = [
        ("OWNER", ["เข้าสู่ระบบด้วยบัญชีเจ้าของร้าน", "ตรวจสอบภาพรวมสินค้า สต็อก ใบเสร็จ และรายการนำส่ง", "จัดการสินค้า หมวดหมู่ หรือสต็อกตามต้องการ"]),
        ("ADMIN", ["เข้าสู่ระบบด้วยบัญชีผู้ดูแลคลัง", "ค้นหาสินค้าที่ต้องการ", "ปรับจำนวนคงเหลือและตรวจสอบผลการเปลี่ยนแปลง"]),
        ("EMPLOYEE", ["เข้าสู่ระบบด้วยบัญชีพนักงาน", "เลือกคำสั่งซื้อที่รับผิดชอบและเปลี่ยนสถานะตามลำดับ", "เมื่ออยู่สถานะ DELIVERING ให้อนุญาต GPS เพื่อแชร์ตำแหน่งระหว่างนำส่ง"]),
        ("RETAILER", ["สมัครหรือเข้าสู่ระบบด้วยบัญชีร้านค้า", "เลือกสินค้า ระบุจำนวนแบบกล่องหรือแพ็ค และตรวจสอบตะกร้า", "ยืนยันคำสั่งซื้อ แล้วติดตามสถานะหรือดูใบเสร็จรายวันของร้านตนเอง"]),
    ]
    for role, steps in instructions:
        add_heading(doc, role, level=2)
        for i, step in enumerate(steps, 1):
            add_text(doc, f"{i}. {step}", first_indent=False, space_after=2)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.page_break_before = True
    p.paragraph_format.space_after = Pt(14)
    set_run_font(p.add_run("ภาคผนวก ข\nแบบบันทึกผลการทดสอบ"), 22, bold=True, color="000000")
    add_text(doc, "ตารางต่อไปนี้ใช้บันทึกผลเมื่อนำระบบไปทดสอบในสภาพแวดล้อมจริง โดยผู้ทดสอบควรระบุวันทดสอบ ผู้ทดสอบ และผลการดำเนินการของแต่ละกรณี", first_indent=False, space_after=8)
    add_table(doc, "ตาราง ข.1 แบบบันทึกผลการทดสอบ", ["กรณี", "ผลที่คาดหวัง", "ผลทดสอบ/หมายเหตุ"], [
        ["เข้าสู่ระบบตามบทบาท", "ไปยังแดชบอร์ดที่ถูกต้อง", ""],
        ["ยืนยันคำสั่งซื้อ", "ข้อมูลออเดอร์ ใบเสร็จ และสต็อกเปลี่ยนแปลงถูกต้อง", ""],
        ["เปลี่ยนสถานะการนำส่ง", "เปลี่ยนได้เฉพาะลำดับที่กำหนด", ""],
        ["แชร์ตำแหน่ง", "บันทึกตำแหน่งล่าสุดได้เมื่ออนุญาต", ""],
    ], [Mm(45), Mm(68), Mm(47)], 9.5)


def set_metadata(doc):
    props = doc.core_properties
    props.title = "การพัฒนาเว็บแอปพลิเคชันระบบจัดการค้าส่ง 123พาณิชย์ปลีกส่ง"
    props.subject = "รายงานโครงงานฉบับปรับให้สอดคล้องกับระบบเว็บ"
    props.author = "ภรัณยู โยธาพล"
    props.comments = "สร้างจากขอบเขตและความสามารถของระบบ 123app"


def main():
    make_assets()
    OUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(16)
    section = doc.sections[0]
    configure_section(section, footer_center=True, number_fmt="lowerRoman", start=1, show_page_number=False)
    set_metadata(doc)
    cover(doc)
    doc.add_section(WD_SECTION.NEW_PAGE)
    preliminary = doc.sections[-1]
    configure_section(preliminary, footer_center=True, number_fmt="lowerRoman", start=1)
    abstract(doc, english=False)
    doc.add_page_break()
    abstract(doc, english=True)
    doc.add_page_break()
    front_matter(doc)
    doc.add_section(WD_SECTION.NEW_PAGE)
    main_section = doc.sections[-1]
    configure_section(main_section, footer_center=False, number_fmt="decimal", start=1)
    chapter1(doc)
    chapter2(doc)
    chapter3(doc)
    chapter4(doc)
    chapter5(doc)
    references_and_appendices(doc)
    doc.save(OUT_DOCX)
    print(OUT_DOCX)


if __name__ == "__main__":
    main()
