#!/usr/bin/env python3
"""Finalize a thesis .docx using Microsoft Word's real pagination (Windows).

  python finalize_thesis.py thesis.docx [--pdf out.pdf]

1. Inserts '(ต่อ)' continuation headers at the top of every page that a
   TOC / list of tables / list of figures / bibliography spills onto
   (page break + title + blank line [+ column header row]).
2. Fills every page number in the TOC and lists from Word's layout,
   repeating until the numbers are stable.
3. Optionally exports a PDF with all fonts embedded.

Works on any docx built by build_thesis.py, and also on hand-edited files
that follow the same conventions (titles 'สารบัญ', 'สารบัญตาราง',
'สารบัญภาพ', 'บรรณานุกรม'; captions 'ภาพที่ x.y ...', 'ตารางที่ x.y ...').
"""
import copy, os, re, subprocess, sys, tempfile
import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
PROBE = os.path.join(HERE, "word_probe.ps1")
LIST_TITLES = ("สารบัญ", "สารบัญตาราง", "สารบัญภาพ")


def roman(n):
    vals = [(1000, "M"), (900, "CM"), (500, "D"), (400, "CD"), (100, "C"), (90, "XC"),
            (50, "L"), (40, "XL"), (10, "X"), (9, "IX"), (5, "V"), (4, "IV"), (1, "I")]
    out = ""
    for v, s in vals:
        while n >= v: out += s; n -= v
    return out


def probe(path):
    out = tempfile.mktemp(suffix=".txt")
    subprocess.run(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", PROBE,
                    os.path.abspath(path), out], check=True, capture_output=True)
    rows = []
    for line in open(out, encoding="utf-8").read().splitlines():
        a, b, t = (line.split("\t", 2) + ["", "", ""])[:3]
        rows.append((int(a or 0), int(b or 0), t))
    os.remove(out)
    return rows


def ptext(p):
    return "".join(t.text or "" for t in p._p.iter(qn("w:t")))


def is_toc_line(p):
    return "\t" in p.text


def body_paras(d):
    return list(d.paragraphs)


def find_title(ps, title):
    for i, p in enumerate(ps):
        if p.text.strip() == title and not is_toc_line(p):
            return i
    return None


def list_block(ps, title):
    """indices of the entries of a TOC/list (after title, blank and the
    'หน้า' header row) up to the next list title / chapter."""
    i = find_title(ps, title)
    if i is None: return None, []
    j = i + 1
    entries = []
    while j < len(ps):
        t = ps[j].text.strip()
        if t in LIST_TITLES and not is_toc_line(ps[j]) and t != title: break
        if re.fullmatch(r"บทที่ \d+", t) and not is_toc_line(ps[j]): break
        if t.startswith(title + " (ต่อ)"): j += 1; continue
        if is_toc_line(ps[j]) and not t.endswith("หน้า"):
            entries.append(j)
        j += 1
    return i, entries


def clone_block(ps, i):
    """title + blank (+ header row) to reuse as the continuation header."""
    els = [copy.deepcopy(ps[i]._p), copy.deepcopy(ps[i + 1]._p)]
    if i + 2 < len(ps) and ps[i + 2].text.strip().endswith("หน้า"):
        els.append(copy.deepcopy(ps[i + 2]._p))
    t = els[0]
    for r in t.iter(qn("w:t")):
        pass
    # rename title -> 'title (ต่อ)'
    ts = list(t.iter(qn("w:t")))
    ts[-1].text = (ts[-1].text or "") + " (ต่อ)"
    pPr = t.find(qn("w:pPr"))
    if pPr.find(qn("w:pageBreakBefore")) is None:
        pPr.insert(0, OxmlElement("w:pageBreakBefore"))
    return els


def remove_continuations(d):
    for p in list(d.paragraphs):
        t = p.text.strip()
        if t.endswith("(ต่อ)") and t[:-5].strip() in LIST_TITLES + ("บรรณานุกรม",):
            nxt = p._p.getnext()
            p._p.getparent().remove(p._p)
            # following blank + header row
            for _ in range(2):
                if nxt is None: break
                tx = "".join(x.text or "" for x in nxt.iter(qn("w:t"))).strip()
                n2 = nxt.getnext()
                if tx == "" or tx.endswith("หน้า"):
                    nxt.getparent().remove(nxt)
                    nxt = n2
                else:
                    break


def reflow(path):
    """insert continuation headers one at a time until stable"""
    for _ in range(40):
        d = docx.Document(path); ps = body_paras(d)
        rows = probe(path)
        assert len(rows) == len(ps), (len(rows), len(ps))
        fixed = False
        # TOC + lists
        for title in LIST_TITLES:
            i, entries = list_block(ps, title)
            if i is None: continue
            if not entries: continue
            cur = rows[i][1]
            for j in range(i + 1, entries[-1] + 1):
                t = ps[j].text.strip()
                if t.startswith(title + " (ต่อ)"):
                    cur = rows[j][1]; continue
                if j not in entries: continue
                if rows[j][1] > cur:
                    # blank 'gap' lines that already moved to the new page
                    # would end up below the continuation header -> drop them
                    k = j - 1
                    while k > i and not ps[k].text.strip() and rows[k][1] > cur:
                        ps[k]._p.getparent().remove(ps[k]._p); k -= 1
                    for e in clone_block(ps, i):
                        ps[j]._p.addprevious(e)
                    fixed = True; break
            if fixed: break
        if not fixed:
            b = find_title(ps, "บรรณานุกรม")
            if b is not None:
                cur = rows[b][1]
                j = b + 1
                while j < len(ps):
                    t = ps[j].text.strip()
                    if t.startswith("ภาคผนวก") or t == "ประวัติผู้จัดทำ": break
                    if t.startswith("บรรณานุกรม (ต่อ)"): cur = rows[j][1]; j += 1; continue
                    if t and rows[j][1] > cur:
                        for e in clone_block(ps, b)[:2]:
                            ps[j]._p.addprevious(e)
                        fixed = True; break
                    j += 1
        if not fixed:
            return
        d.save(path)
    raise RuntimeError("continuation reflow did not converge")


def keep_bib_entries_whole(path):
    d = docx.Document(path); ps = body_paras(d)
    b = find_title(ps, "บรรณานุกรม")
    if b is None: return
    for p in ps[b + 1:]:
        t = p.text.strip()
        if t.startswith("ภาคผนวก") or t == "ประวัติผู้จัดทำ": break
        if t:
            pPr = p._p.get_or_add_pPr()
            if pPr.find(qn("w:keepLines")) is None:
                pPr.insert(0, OxmlElement("w:keepLines"))
    d.save(path)


def set_number(p, value):
    els = list(p._p.iter(qn("w:tab"), qn("w:t")))
    last_tab = max(i for i, e in enumerate(els) if e.tag == qn("w:tab"))
    after = [e for e in els[last_tab + 1:] if e.tag == qn("w:t")]
    if not after: return False
    if after[0].text == value and all(not e.text for e in after[1:]): return False
    after[0].text = value
    for e in after[1:]: e.text = ""
    return True


def fill_numbers(path):
    d = docx.Document(path); ps = body_paras(d)
    rows = probe(path)
    assert len(rows) == len(ps)
    front_end = next(i for i, p in enumerate(ps) if re.fullmatch(r"บทที่ \d+", p.text.strip()))

    def disp(i):
        return roman(rows[i][0]) if i < front_end else str(rows[i][0])

    targets = {}
    for i, p in enumerate(ps):
        t = p.text.strip()
        if not t or is_toc_line(p): continue
        if i < front_end and t in ("บทคัดย่อ", "กิตติกรรมประกาศ") + LIST_TITLES:
            # the abstract heading sits below the info block; page is what counts
            targets.setdefault(t, disp(i))
        if re.fullmatch(r"บทที่ \d+", t):
            targets[t] = disp(i)
        elif re.match(r"^\d+\.\d+ \S", t):
            targets.setdefault(t, disp(i))
        elif re.match(r"^(ภาพที่|ตารางที่) (\d+|[ก-ฮ])\.\d+ ", t):
            targets.setdefault(re.match(r"^(ภาพที่|ตารางที่) (\S+)", t).group(0), disp(i))
        elif t in ("บรรณานุกรม", "ภาคผนวก", "ประวัติผู้จัดทำ"):
            targets.setdefault(t, disp(i))
        elif re.fullmatch(r"ภาคผนวก [ก-ฮ]", t):
            targets.setdefault(t, disp(i))
    # the TOC itself is the first 'สารบัญ'
    changed = 0
    mode = None; prev_key = None
    for i, p in enumerate(ps[:front_end]):
        t = p.text.strip()
        if not is_toc_line(p):
            base = t.replace(" (ต่อ)", "")
            if base == "สารบัญภาพ": mode = "ภาพที่"
            elif base == "สารบัญตาราง": mode = "ตารางที่"
            elif base == "สารบัญ": mode = "toc"
            continue
        if t.endswith("หน้า") or mode is None: continue
        key, _, num = p.text.rpartition("\t"); key = key.strip()
        if mode == "toc":
            if key.startswith("บทที่"):
                tg = targets.get(" ".join(key.split()[:2]))
            elif re.match(r"^ภาคผนวก [ก-ฮ] ", key):
                tg = targets.get(" ".join(key.split()[:2]))
            else:
                tg = targets.get(key)
        else:
            m = re.match(r"^(\d+\.\d+|[ก-ฮ]\.\d+)\s", key)
            tg = targets.get(f"{mode} {m.group(1)}") if m else None
        if tg is None:
            print("  WARNING: no page found for", key); continue
        if set_number(p, tg): changed += 1
    d.save(path)
    return changed


THAI = re.compile(r"[฀-๿]")


def stretched_lines(pdf):
    """Thai lines whose letters were spread apart by Thai-distributed
    justification (happens when a line cannot break at a word boundary)."""
    import fitz
    out = []
    for p in fitz.open(pdf):
        for b in p.get_text("rawdict")["blocks"]:
            for l in b.get("lines", []):
                ch = [c for s in l["spans"] for c in s["chars"]]
                tx = "".join(c["c"] for c in ch)
                if len(ch) < 20 or not THAI.search(tx): continue
                gaps = [q["bbox"][0] - a["bbox"][2] for a, q in zip(ch, ch[1:])
                        if a["c"] != " " and q["c"] != " " and re.match("[ก-ฮ]", q["c"])]
                if gaps and sum(1 for g in gaps if g > 1.2) / len(gaps) > 0.35:
                    out.append(tx)
    return out


def fix_stretched(path, rounds=3):
    """Condense character spacing slightly (0.3 pt, then 0.6 pt) in paragraphs
    whose lines render letter-spaced, so the next word fits and the line is
    justified at word boundaries instead - the same manual fix used on the
    approved thesis."""
    norm = lambda t: re.sub(r"\s+", "", t.replace("​", ""))
    for r in range(rounds):
        tmp = tempfile.mktemp(suffix=".pdf")
        export_pdf(path, tmp)
        bad = stretched_lines(tmp); os.remove(tmp)
        if not bad: return
        d = docx.Document(path); n = 0
        for line in bad:
            key = norm(line)[:20]
            for p in d.paragraphs:
                if key and key in norm(p.text):
                    for rr in p._p.iter(qn("w:r")):
                        rPr = rr.find(qn("w:rPr"))
                        if rPr is None: continue
                        sp = rPr.find(qn("w:spacing"))
                        if sp is None:
                            sp = OxmlElement("w:spacing"); rPr.append(sp); sp.set(qn("w:val"), "0")
                        sp.set(qn("w:val"), str(max(-12, int(sp.get(qn("w:val"))) - 6)))
                    n += 1; break
        d.save(path)
        print(f"  condensed {n} paragraph(s) with stretched lines")


def export_pdf(docx_path, pdf_path):
    ps = f"""
$w=New-Object -ComObject Word.Application; $w.Visible=$false; $w.DisplayAlerts=0
try {{
 $d=$w.Documents.Open('{os.path.abspath(docx_path)}',$false,$true)
 $d.Repaginate()
 $d.ExportAsFixedFormat('{os.path.abspath(pdf_path)}',17,$false,0,0,0,0,0,$true,$true,1,$true,$true,$false)
 $d.Close($false)
}} finally {{ $w.Quit() }}
"""
    f = tempfile.mktemp(suffix=".ps1")
    open(f, "w", encoding="utf-8-sig").write(ps)
    subprocess.run(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", f], check=True)
    os.remove(f)


def main():
    path = sys.argv[1]
    pdf = None
    if "--pdf" in sys.argv:
        pdf = sys.argv[sys.argv.index("--pdf") + 1]
    keep_bib_entries_whole(path)
    fix_stretched(path)
    for it in range(6):
        reflow(path)
        n = fill_numbers(path)
        print(f"pass {it + 1}: {n} page numbers changed")
        if n == 0:
            break
    if pdf:
        export_pdf(path, pdf)
        print("pdf", pdf)


if __name__ == "__main__":
    main()
