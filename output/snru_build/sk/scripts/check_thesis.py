#!/usr/bin/env python3
"""Audit a thesis docx (+ its PDF) against the SNRU format rules.

  python check_thesis.py thesis.docx thesis.pdf

Prints one line per rule: OK / FAIL with details. Exit code 1 if any FAIL.
The rules are the ones in references/format-spec.md; every rule here was a
real problem found while preparing an actual thesis, so treat a FAIL as a
real defect, not a style preference.
"""
import re, sys, os
import docx
from docx.oxml.ns import qn

THAI = re.compile(r"[฀-๿]")
FONT_OK = ("TH SarabunPSK", "TH Sarabun PSK")
results = []


def rule(name, ok, detail=""):
    results.append((ok, name, detail))


def ptext(p):
    return p.text


def runs(p):
    for r in p._p.iter(qn("w:r")):
        if r.getparent().tag.endswith("}oMath") or r.getparent().tag.endswith("}f") : continue
        t = "".join(x.text or "" for x in r.iter(qn("w:t")))
        yield r, t


STY = {}


def init_styles(d):
    st = d.styles.element
    STY["styles"] = {e.get(qn("w:styleId")): e for e in st.findall(qn("w:style"))}
    STY["default"] = st.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))


def _chain(sid):
    out = []
    while sid and sid in STY["styles"]:
        s = STY["styles"][sid]; out.append(s)
        b = s.find(qn("w:basedOn")); sid = b.get(qn("w:val")) if b is not None else None
    return out


def val(rPr, tag, r=None):
    """effective run property: direct -> char style -> paragraph style chain -> defaults"""
    if rPr is not None and rPr.find(qn(tag)) is not None:
        return rPr.find(qn(tag)).get(qn("w:val"), "1")
    if r is None:
        return None
    chain = []
    if rPr is not None and rPr.find(qn("w:rStyle")) is not None:
        chain += _chain(rPr.find(qn("w:rStyle")).get(qn("w:val")))
    p = r.getparent()
    while p is not None and p.tag != qn("w:p"):
        p = p.getparent()
    pPr = p.find(qn("w:pPr")) if p is not None else None
    ps = pPr.find(qn("w:pStyle")) if pPr is not None else None
    chain += _chain(ps.get(qn("w:val")) if ps is not None else "Normal")
    for s in chain:
        x = s.find(qn("w:rPr"))
        if x is not None and x.find(qn(tag)) is not None:
            return x.find(qn(tag)).get(qn("w:val"), "1")
    e = STY["default"].find(qn(tag)) if STY.get("default") is not None else None
    return e.get(qn("w:val"), "1") if e is not None else None


def on(v):
    return v is not None and v not in ("0", "false")


def check_docx(path):
    d = docx.Document(path)
    init_styles(d)
    ps = d.paragraphs
    # margins
    bad = []
    for i, s in enumerate(d.sections):
        pm = s._sectPr.find(qn("w:pgMar"))
        got = tuple(pm.get(qn("w:" + k)) for k in ("top", "left", "right", "bottom"))
        if got != ("2160", "2160", "1440", "1440"): bad.append((i, got))
    rule("margins top/left 1.5in, right/bottom 1in (manual 4.6)", not bad, str(bad[:3]))
    # fonts + paired bold/size + spacing
    fonts, pair, spacing = set(), [], []
    for p in ps + [c for t in d.tables for row in t.rows for cell in row.cells for c in cell.paragraphs]:
        pPr = p._p.pPr
        sp = pPr.find(qn("w:spacing")) if pPr is not None else None
        if sp is not None and (sp.get(qn("w:before"), "0") != "0" or sp.get(qn("w:after"), "0") != "0"):
            spacing.append(p.text[:30])
        for r, t in runs(p):
            if not t.strip(): continue
            rPr = r.find(qn("w:rPr"))
            f = rPr.find(qn("w:rFonts")) if rPr is not None else None
            for a in ("ascii", "hAnsi", "cs"):
                v = f.get(qn("w:" + a)) if f is not None else None
                if v: fonts.add(v)
    rule("all text in TH SarabunPSK (manual 4.2.1)", fonts <= set(FONT_OK) | {"Times New Roman"}, str(sorted(fonts)))
    rule("paragraph spacing before/after = 0 (gaps are Enter only)", not spacing, f"{len(spacing)} e.g. {spacing[:3]}")
    # glyph coverage
    missing = set()
    try:
        from fontTools.ttLib import TTFont
        cand = [os.path.expandvars(r"%LOCALAPPDATA%\Microsoft\Windows\Fonts\THSarabunPSK Regular.ttf"),
                r"C:\Windows\Fonts\THSarabunPSK.ttf", r"C:\Windows\Fonts\THSarabunNew.ttf"]
        fp = next((c for c in cand if os.path.exists(c)), None)
        if fp:
            cmap = TTFont(fp).getBestCmap()
            for p in ps:
                for r, t in runs(p):
                    rPr = r.find(qn("w:rPr")); f = rPr.find(qn("w:rFonts")) if rPr is not None else None
                    if f is not None and f.get(qn("w:ascii")) == "Times New Roman":
                        continue   # deliberately set to Times New Roman (symbols / math)
                    for ch in t:
                        if ord(ch) > 32 and ch not in "\u200b\t" and ord(ch) not in cmap: missing.add(ch)
    except Exception:
        pass
    rule("no characters missing from TH SarabunPSK (they fall back to another font)", not missing,
         " ".join(f"U+{ord(c):04X}" for c in sorted(missing)))
    # double blank lines (outside covers / vertically centred dividers)
    body = list(d.element.body.iterchildren())
    def blank(e):
        return e.tag == qn("w:p") and not "".join(x.text or "" for x in e.iter(qn("w:t"))).strip() \
            and not any(k in e.xml for k in ("a:blip", "m:oMath", "w:drawing", "w:sectPr", "w:br "))
    start = next(i for i, e in enumerate(body) if e.tag == qn("w:p") and
                 "".join(x.text or "" for x in e.iter(qn("w:t"))).strip() == "บทคัดย่อ") if any(
        p.text.strip() == "บทคัดย่อ" for p in ps) else 0
    dbl = []
    k = start
    while k < len(body):
        if blank(body[k]):
            j = k
            while j < len(body) and blank(body[j]): j += 1
            nxt = body[j] if j < len(body) else None
            prev = body[k - 1]
            at_break = nxt is None or "w:sectPr" in nxt.xml or "pageBreakBefore" in nxt.xml or nxt.tag == qn("w:sectPr")
            # allowed: top of a new section (manual vertical positioning of a
            # divider page) and the space reserved under a floating picture
            exempt = "w:sectPr" in prev.xml or "wp:anchor" in prev.xml
            if j - k >= 2 and not at_break and not exempt:
                what = "".join(x.text or "" for x in prev.iter(qn("w:t")))[:25] or ("[table]" if prev.tag == qn("w:tbl") else "[picture/equation]")
                dbl.append(f"after '{what}' ({j-k} blank lines)")
            k = j
        else:
            k += 1
    rule("never two blank lines in a row after the covers", not dbl, str(dbl[:5]))
    # x.y headings must keep with the next paragraph (direct or via style)
    def keep_next(p):
        pPr = p._p.pPr
        if pPr is not None and pPr.find(qn("w:keepNext")) is not None:
            return pPr.find(qn("w:keepNext")).get(qn("w:val"), "1") not in ("0", "false")
        ps_ = pPr.find(qn("w:pStyle")) if pPr is not None else None
        for st in _chain(ps_.get(qn("w:val")) if ps_ is not None else "Normal"):
            x = st.find(qn("w:pPr"))
            if x is not None and x.find(qn("w:keepNext")) is not None:
                return x.find(qn("w:keepNext")).get(qn("w:val"), "1") not in ("0", "false")
        return False
    nokeep = [p.text[:25] for p in ps if re.match(r"^\d+\.\d+(\.\d+)? \S", p.text.strip()) and len(p.text) < 90
              and "	" not in p.text and p.runs and all((r.bold or False) for r in p.runs if r.text.strip())
              and not keep_next(p)]
    rule("headings are kept with the following paragraph", not nokeep, str(nokeep[:5]))
    # table shading (white fill hides borders in PDF)
    shd = sum(1 for t in d.tables for s in t._tbl.iter(qn("w:shd")) if (s.get(qn("w:fill")) or "").upper() == "FFFFFF")
    rule("no white cell shading in tables (it hides borders in the PDF)", shd == 0, f"{shd} cells")
    return d


def check_pdf(pdf, d):
    import fitz
    doc = fitz.open(pdf)
    emb = [f[3] for p in doc for f in p.get_fonts() if f[1] in ("n/a", "")]
    rule("all PDF fonts embedded", not emb, str(emb))
    other = set()
    for i, p in enumerate(doc):
        for b in p.get_text("rawdict")["blocks"]:
            for l in b.get("lines", []):
                for s in l["spans"]:
                    t = "".join(c["c"] for c in s["chars"])
                    if t.strip() and "Sarabun" not in s["font"] and "TimesNewRoman" not in s["font"].replace(" ", ""):
                        other.add((i + 1, s["font"], t[:10]))
    rule("every visible character rendered in TH SarabunPSK / Times (math)", not other, str(sorted(other)[:5]))
    # rendered sizes / bold of each role (ground truth = what the PDF shows)
    issues = []
    started = False
    for i, p in enumerate(doc):
        for b in p.get_text("rawdict")["blocks"]:
            for l in b.get("lines", []):
                sp = [s for s in l["spans"] if "".join(c["c"] for c in s["chars"]).strip()]
                if not sp: continue
                tx = "".join("".join(c["c"] for c in s["chars"]) for s in sp).strip()
                sizes = {round(s["size"]) for s in sp}; bolds = ["Bold" in s["font"] for s in sp]
                if re.fullmatch(r"บทที่ \d+", tx) and l["bbox"][1] < 130:
                    started = True
                    if sizes != {20} or not all(bolds): issues.append(f"p{i+1} chapter {tx}: {sizes}")
                    continue
                if not started: continue
                if re.match(r"^\d+\.\d+ \S", tx) and all(bolds) and len(tx) < 90:
                    if sizes != {18}: issues.append(f"p{i+1} x.y heading 18: {tx[:20]} {sizes}")
                elif re.match(r"^\d+\.\d+\.\d+ \S", tx) and bolds[0] and len(tx) < 110:
                    if sizes != {16} or not all(bolds): issues.append(f"p{i+1} x.y.z heading 16 bold: {tx[:20]}")
                elif re.match(r"^(ภาพที่|ตารางที่) (\d+|[ก-ฮ])\.\d+ ", tx) and bolds[0]:
                    lab = re.match(r"^((ภาพที่|ตารางที่) \S+)", tx).group(1)
                    acc = ""; ok = True
                    for s in sp:
                        st = "".join(c["c"] for c in s["chars"])
                        inlab = len(acc.strip()) < len(lab)
                        if st.strip() and inlab and "Bold" not in s["font"]: ok = False
                        if st.strip() and not inlab and "Bold" in s["font"]: ok = False
                        acc += st
                    if not ok or sizes != {16}: issues.append(f"p{i+1} caption: {tx[:25]}")
                elif tx.startswith("ที่มา"):
                    if any(bolds): issues.append(f"p{i+1} source bold: {tx[:20]}")
    rule("rendered heading / caption sizes and bold follow the manual (4.8-4.10)", not issues, "; ".join(issues[:6]))
    # page numbers
    W, H = doc[0].rect.width, doc[0].rect.height
    nonum_ok, wrong = [], []
    for i, p in enumerate(doc):
        blocks = [b for b in p.get_text("blocks") if b[4].strip()]
        top = [b for b in blocks if b[1] < 80 and re.fullmatch(r"\d+", b[4].strip())]
        bottom = [b for b in blocks if b[1] > H - 70 and re.fullmatch(r"[IVXLC]+", b[4].strip())]
        first = next((b[4].strip() for b in blocks if b not in top + bottom), "")
        must_hide = bool(re.match(r"บทที่ \d+", first)) or first.startswith("บรรณานุกรม\n") or first == "บรรณานุกรม" \
            or first.startswith("ภาคผนวก")
        if must_hide and top:
            wrong.append(f"p{i+1} '{first[:12]}' shows {top[0][4].strip()}")
    rule("first page of every chapter / bibliography / appendix shows no page number (4.7.3)", not wrong, "; ".join(wrong[:5]))
    # near-empty pages
    empty = []
    for i, p in enumerate(doc):
        t = re.sub(r"\s+", " ", p.get_text()).strip()
        if len(t) < 15 and not p.get_images() and not t.startswith("ภาคผนวก") and t != "บรรณานุกรม":
            empty.append(i + 1)
    rule("no blank / near-empty pages", not empty, str(empty))
    # text inside margins
    out = []
    for i, p in enumerate(doc):
        for b in p.get_text("blocks"):
            if b[4].strip() and (b[2] > W - 60 or b[0] < 95 or b[3] > H - 60) and not re.fullmatch(r"[0-9IVXLC]+", b[4].strip()):
                out.append(f"p{i+1}:{b[4].strip()[:12]}")
    rule("no text outside the margins", not out, str(out[:5]))
    # stretched (letter-spaced) Thai lines
    spaced = []
    for i, p in enumerate(doc):
        for b in p.get_text("rawdict")["blocks"]:
            for l in b.get("lines", []):
                ch = [c for s in l["spans"] for c in s["chars"]]
                tx = "".join(c["c"] for c in ch)
                if len(ch) < 20 or not THAI.search(tx): continue
                gaps = [q["bbox"][0] - a["bbox"][2] for a, q in zip(ch, ch[1:])
                        if a["c"] != " " and q["c"] != " " and re.match("[ก-ฮ]", q["c"])]
                if gaps and sum(1 for g in gaps if g > 1.2) / len(gaps) > 0.35:
                    spaced.append(f"p{i+1}:{tx[:15]}")
    rule("no letter-spaced (over-stretched) Thai lines", len(spaced) <= 2, f"{len(spaced)}: {spaced[:4]}")
    # heading stranded at page bottom
    strand = []
    def lines_of(pg):
        out = []
        for b in pg.get_text("rawdict")["blocks"]:
            for l in b.get("lines", []):
                sp = [s for s in l["spans"] if "".join(c["c"] for c in s["chars"]).strip()]
                tx = "".join("".join(c["c"] for c in s["chars"]) for s in sp).strip()
                if sp and not re.fullmatch(r"[0-9IVX]+", tx):
                    out.append((l["bbox"][1], tx, all("Bold" in s["font"] for s in sp)))
        return sorted(out)
    for i, p in enumerate(doc):
        if i <= 8: continue
        ls = lines_of(p)
        if not ls: continue
        y, tx, bold = ls[-1]
        if bold and re.match(r"^\d+\.\d+(\.\d+)? \S", tx) and len(tx) < 90:
            nxt = lines_of(doc[i + 1]) if i + 1 < len(doc) else []
            nf = nxt[0][1] if nxt else ""
            depth = tx.split()[0].count(".")
            ends_chapter = re.match(r"^(บทที่ \d|บรรณานุกรม|ภาคผนวก)", nf)   # last item of a chapter
            if not ends_chapter and not (re.match(r"^\d+(\.\d+)+ ", nf) and nf.split()[0].count(".") == depth):
                strand.append(f"p{i+1}:{tx[:20]}")
    rule("no heading left alone at the bottom of a page", not strand, str(strand))
    return doc


def main():
    d = check_docx(sys.argv[1])
    if len(sys.argv) > 2:
        check_pdf(sys.argv[2], d)
    fails = 0
    for ok, name, det in results:
        print(("OK   " if ok else "FAIL ") + name + ("" if ok else "  -> " + det))
        fails += (not ok)
    print(f"{len(results) - fails}/{len(results)} rules passed")
    sys.exit(1 if fails else 0)


if __name__ == "__main__":
    main()
