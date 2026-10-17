#!/usr/bin/env python3
"""Build an SNRU computer-project report (.docx) from a content JSON file.

Usage:
    python build_thesis.py content.json out.docx

The JSON schema is documented in references/content-schema.md. Formatting
follows references/format-spec.md (SNRU manual + advisor-approved layout).
Page numbers in the TOC / lists are written as placeholders; run
finalize_thesis.py afterwards to fill them in from Word's real pagination.
"""
import copy, json, os, re, sys
import docx
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Emu, Cm

FONT = "TH SarabunPSK"
HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(os.path.dirname(HERE), "assets")
TEXT_W = 8306          # A4 width 11906 - left 2160 - right 1440 (twips)
THAI = re.compile(r"[฀-๿]")
M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"


def mel(tag):
    return OxmlElement("m:" + tag)


# ----------------------------------------------------------------- low level
_CMAP = None


def sarabun_cmap():
    global _CMAP
    if _CMAP is None:
        _CMAP = set()
        try:
            from fontTools.ttLib import TTFont
            for c in (os.path.expandvars(r"%LOCALAPPDATA%\Microsoft\Windows\Fonts\THSarabunPSK Regular.ttf"),
                      r"C:\Windows\Fonts\THSarabunPSK.ttf", r"C:\Windows\Fonts\THSarabun.ttf"):
                if os.path.exists(c):
                    _CMAP = set(TTFont(c).getBestCmap()); break
        except Exception:
            pass
    return _CMAP


def set_rfonts(rPr, font=FONT):
    f = rPr.find(qn("w:rFonts"))
    if f is None:
        f = OxmlElement("w:rFonts"); rPr.insert(0, f)
    for a in ("ascii", "hAnsi", "cs", "eastAsia"):
        f.set(qn("w:" + a), font)


def add_run(p, text, bold=False, italic=False, size=16, font=FONT):
    """Add text as runs; Thai and non-Thai segments get separate runs so the
    complex-script flag (w:cs) is only on Thai.  Bold/italic/size are always
    written to BOTH the Latin and complex-script slots (b+bCs, i+iCs,
    sz+szCs) - Word renders Thai from the cs slots, and a mismatch is what
    makes Thai text silently turn bold/regular or change size."""
    cmap = sarabun_cmap()
    segs = []
    for seg in re.findall(r"[฀-๿][฀-๿\s]*|[^฀-๿]+", text):
        # characters TH SarabunPSK does not contain (Greek letters, arrows...)
        # get Times New Roman explicitly; otherwise Word silently picks some
        # random fallback font (Calibri, Angsana...)
        if cmap and font == FONT:
            buf = ""; cur = None
            for ch in seg:
                miss = ord(ch) > 32 and ord(ch) not in cmap and ch != "​"
                f_ = "Times New Roman" if miss else FONT
                if cur is not None and f_ != cur:
                    segs.append((buf, cur)); buf = ""
                buf += ch; cur = f_
            if buf: segs.append((buf, cur))
        else:
            segs.append((seg, font))
    for seg, sfont in segs:
        r = OxmlElement("w:r"); rPr = OxmlElement("w:rPr"); r.append(rPr)
        set_rfonts(rPr, sfont)
        if bold:
            rPr.append(OxmlElement("w:b")); rPr.append(OxmlElement("w:bCs"))
        if italic:
            rPr.append(OxmlElement("w:i")); rPr.append(OxmlElement("w:iCs"))
        for tag in ("w:sz", "w:szCs"):
            e = OxmlElement(tag); e.set(qn("w:val"), str(int(size * 2))); rPr.append(e)
        if THAI.search(seg):
            rPr.append(OxmlElement("w:cs"))
        lang = OxmlElement("w:lang"); lang.set(qn("w:bidi"), "th-TH"); rPr.append(lang)
        t = OxmlElement("w:t"); t.set(XML_SPACE, "preserve"); t.text = seg; r.append(t)
        p._p.append(r)


def add_marked(p, text, size=16, bold=False):
    """**bold**, *italic* inline markup."""
    for part in re.split(r"(\*\*.+?\*\*|\*[^*]+?\*)", text):
        if not part:
            continue
        if part.startswith("**"):
            add_run(p, part[2:-2], bold=True, size=size)
        elif part.startswith("*") and len(part) > 1:
            add_run(p, part[1:-1], italic=True, bold=bold, size=size)
        else:
            add_run(p, part, bold=bold, size=size)


def add_tab(p, bold=False, size=16):
    r = OxmlElement("w:r"); rPr = OxmlElement("w:rPr"); r.append(rPr); set_rfonts(rPr)
    if bold:
        rPr.append(OxmlElement("w:b")); rPr.append(OxmlElement("w:bCs"))
    for tag in ("w:sz", "w:szCs"):
        e = OxmlElement(tag); e.set(qn("w:val"), str(size * 2)); rPr.append(e)
    r.append(OxmlElement("w:tab")); p._p.append(r)


def fmt(p, align="thai", first=0, left=0, hanging=0, keep_next=False,
        page_break=False, tabs=None, size=16, keep_lines=False, outline=None):
    """Paragraph formatting. Spacing before/after is always 0 and line
    spacing single: vertical gaps are made only with empty paragraphs."""
    pPr = p._p.get_or_add_pPr()
    for tag in ("w:keepNext", "w:keepLines", "w:pageBreakBefore"):
        e = pPr.find(qn(tag))
        if e is not None:
            pPr.remove(e)
    if keep_next: pPr.append(OxmlElement("w:keepNext"))
    if keep_lines: pPr.append(OxmlElement("w:keepLines"))
    if page_break: pPr.append(OxmlElement("w:pageBreakBefore"))
    if tabs:
        tb = OxmlElement("w:tabs")
        for (val, pos, leader) in tabs:
            t = OxmlElement("w:tab"); t.set(qn("w:val"), val); t.set(qn("w:pos"), str(pos))
            if leader: t.set(qn("w:leader"), leader)
            tb.append(t)
        pPr.append(tb)
    sp = OxmlElement("w:spacing")
    sp.set(qn("w:before"), "0"); sp.set(qn("w:after"), "0")
    sp.set(qn("w:line"), "240"); sp.set(qn("w:lineRule"), "auto"); pPr.append(sp)
    ind = OxmlElement("w:ind")
    if left: ind.set(qn("w:left"), str(left))
    if first: ind.set(qn("w:firstLine"), str(first))
    if hanging: ind.set(qn("w:hanging"), str(hanging))
    pPr.append(ind)
    jc = OxmlElement("w:jc")
    jc.set(qn("w:val"), {"thai": "thaiDistribute", "left": "left", "center": "center",
                         "right": "right", "both": "both"}[align])
    pPr.append(jc)
    if outline is not None:
        o = OxmlElement("w:outlineLvl"); o.set(qn("w:val"), str(outline)); pPr.append(o)
    rpr = OxmlElement("w:rPr"); set_rfonts(rpr)
    for tag in ("w:sz", "w:szCs"):
        e = OxmlElement(tag); e.set(qn("w:val"), str(size * 2)); rpr.append(e)
    pPr.append(rpr)
    return p


# --------------------------------------------------------------- the builder
class Builder:
    def __init__(self, data, base_dir):
        self.d = docx.Document()
        self.data = data
        self.base = base_dir
        self.fig = {}          # chapter -> counter
        self.tab = {}
        self.toc = []          # (kind, text, level)
        self.lof = []; self.lot = []
        self.first_section = True
        self._setup_document()

    # -- document wide defaults
    def _setup_document(self):
        st = self.d.styles
        normal = st["Normal"]; normal.font.name = FONT; normal.font.size = Pt(16)
        rPr = normal.element.get_or_add_rPr(); set_rfonts(rPr)
        for tag in ("w:sz", "w:szCs"):
            e = rPr.find(qn(tag))
            if e is None: e = OxmlElement(tag); rPr.append(e)
            e.set(qn("w:val"), "32")
        lang = rPr.find(qn("w:lang"))
        if lang is None: lang = OxmlElement("w:lang"); rPr.append(lang)
        lang.set(qn("w:val"), "en-US"); lang.set(qn("w:bidi"), "th-TH")
        pf = normal.paragraph_format; pf.space_before = Pt(0); pf.space_after = Pt(0); pf.line_spacing = 1.0
        # document defaults too (theme fonts otherwise leak into page numbers, spaces...)
        dd = st.element.find(qn("w:docDefaults"))
        rd = dd.find(qn("w:rPrDefault")).find(qn("w:rPr"))
        f = rd.find(qn("w:rFonts"))
        if f is not None: rd.remove(f)
        set_rfonts(rd)
        lg = rd.find(qn("w:lang"))
        if lg is None: lg = OxmlElement("w:lang"); rd.append(lg)
        lg.set(qn("w:bidi"), "th-TH")
        sec = self.d.sections[0]
        self._page_setup(sec)
        # remove the empty paragraph python-docx starts with
        body = self.d.element.body
        for p in list(body.findall(qn("w:p"))):
            body.remove(p)

    def _page_setup(self, sec):
        sec.page_width = Emu(7560310); sec.page_height = Emu(10692130)
        pm = sec._sectPr.find(qn("w:pgMar"))
        for k, v in dict(top=2160, right=1440, bottom=1440, left=2160, header=1020, footer=720, gutter=0).items():
            pm.set(qn("w:" + k), str(v))

    def p(self, text="", **kw):
        par = self.d.add_paragraph()
        bold = kw.pop("bold", False); size = kw.get("size", 16)
        fmt(par, **kw)
        if text:
            add_marked(par, text, size=size, bold=bold)
        return par

    def blank(self, size=16, keep_next=False):
        """one empty line (= pressing Enter once). Never stacks: if the
        previous paragraph is already blank nothing is added, because two
        blank lines in a row is a formatting error in this layout."""
        body = self.d.element.body
        prev = body[-2] if len(body) > 1 and body[-1].tag == qn("w:sectPr") else (body[-1] if len(body) else None)
        if prev is not None and prev.tag == qn("w:p") and not "".join(t.text or "" for t in prev.iter(qn("w:t"))).strip()                 and not any(k in prev.xml for k in ("w:drawing", "m:oMath", "w:sectPr")) and "lineRule=\"exact\"" not in prev.xml:
            return None
        return self.p("", align="left", size=size, keep_next=keep_next)

    # -- sections and page numbers
    def new_section(self, number=None, restart=None, hide_first=True, vcenter=False):
        """Close the current section and start a new one on a new page.
        number: None (no page numbers), 'roman' (footer centre), 'arabic'
        (header right)."""
        if self.first_section:
            sec = self.d.sections[0]; self.first_section = False
        else:
            sec = self.d.add_section(WD_SECTION.NEW_PAGE)
        self._page_setup(sec)
        sp = sec._sectPr
        for tag in ("w:pgNumType", "w:vAlign", "w:titlePg"):
            for e in sp.findall(qn(tag)):
                sp.remove(e)
        if restart is not None or number == "roman":
            pg = OxmlElement("w:pgNumType")
            if number == "roman": pg.set(qn("w:fmt"), "upperRoman")
            if restart is not None: pg.set(qn("w:start"), str(restart))
            cols = sp.find(qn("w:cols"))
            (cols.addprevious(pg) if cols is not None else sp.append(pg))
        if vcenter:
            v = OxmlElement("w:vAlign"); v.set(qn("w:val"), "center")
            dg = sp.find(qn("w:docGrid"))
            (dg.addprevious(v) if dg is not None else sp.append(v))
        sec.different_first_page_header_footer = bool(hide_first and number)
        for part in (sec.header, sec.footer, sec.first_page_header, sec.first_page_footer):
            part.is_linked_to_previous = False
            for p in part.paragraphs:
                for r in list(p._p):
                    if r.tag != qn("w:pPr"): p._p.remove(r)
        if number == "roman":
            self._page_field(sec.footer.paragraphs[0], "center")
        elif number == "arabic":
            self._page_field(sec.header.paragraphs[0], "right")
        return sec

    def _page_field(self, par, align):
        fmt(par, align=align)
        def r(child_tag=None, text=None, fld=None):
            run = OxmlElement("w:r"); rPr = OxmlElement("w:rPr"); set_rfonts(rPr)
            for tag in ("w:sz", "w:szCs"):
                e = OxmlElement(tag); e.set(qn("w:val"), "32"); rPr.append(e)
            run.append(rPr)
            if fld:
                fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), fld); run.append(fc)
            if child_tag == "instr":
                it = OxmlElement("w:instrText"); it.set(XML_SPACE, "preserve"); it.text = " PAGE "; run.append(it)
            if text is not None:
                t = OxmlElement("w:t"); t.text = text; run.append(t)
            par._p.append(run)
        r(fld="begin"); r("instr"); r(fld="separate"); r(text="1"); r(fld="end")

    # -- images
    def image(self, path, width_cm=None, height_cm=None, align="center"):
        par = self.p("", align=align, keep_next=True)
        run = par.add_run()
        kw = {}
        if width_cm: kw["width"] = Cm(min(width_cm, 14.6))   # text width is 14.65 cm
        if height_cm: kw["height"] = Cm(height_cm)
        run.add_picture(path, **kw)
        return par

    def path(self, p):
        return p if os.path.isabs(p) else os.path.join(self.base, p)

    # ------------------------------------------------------------ front matter
    # Line heights Word uses for TH SarabunPSK at single spacing (points),
    # measured from the approved thesis PDF. Used to place cover elements at
    # the exact vertical positions of the approved layout.
    LH = {16: 21.0, 18: 23.6, 20: 26.3, 22: 28.8}
    LOGO_LINE = 85.5          # height of the logo line (logo is 2.36 cm wide)

    def spacer(self, pts):
        """empty paragraph with an exact line height (spacing before/after
        stay 0) - used only on cover pages for precise vertical placement"""
        if pts < 0.5: return
        par = self.d.add_paragraph()
        fmt(par, align="left")
        sp = par._p.pPr.find(qn("w:spacing"))
        sp.set(qn("w:line"), str(int(round(pts * 20)))); sp.set(qn("w:lineRule"), "exact")

    def _cover_page(self, logo, blocks, first):
        """blocks: list of (target_y or None, [(text,size,bold,align,left)...])
        y positions are measured from the top edge of the page."""
        par = self.image(logo, width_cm=2.36)
        fmt(par, align="center", page_break=not first, keep_next=True)
        y = 108 + self.LOGO_LINE
        for target, lines in blocks:
            if target is not None:
                self.spacer(target - y); y = max(y, target)
            for (text, size, bold, align, left) in lines:
                if text is None:   # cover pages may need two blank lines in a row
                    self.p("", align="left", size=size)
                elif align == "labelline":   # "อาจารย์ที่ปรึกษา ......... ____" on one row
                    par = self.p("", align="left", size=size, tabs=[("left", left, None)])
                    add_run(par, text[0], size=size); add_tab(par, size=size); add_run(par, text[1], size=size)
                else:
                    self.p(text, align=align, size=size, bold=bold, left=left)
                y += self.LH[size]

    def covers(self):
        m = self.data["meta"]
        logo = m.get("logo") and self.path(m["logo"]) or os.path.join(ASSETS, "snru_logo.png")
        self.new_section(number=None)
        degree_line = f"ปริญญา{m['degree_th']} สาขาวิชา{m['program_th']}"
        foot = [m.get("cover_statement", "โครงงานคอมพิวเตอร์นี้เป็นส่วนหนึ่งของการศึกษาตามหลักสูตร"),
                degree_line, m["university_th"], f"พ.ศ. {m['year_be']}"]
        C = lambda t, sz, b=True: (t, sz, b, "center", 0)
        # 1 outer cover (Thai + English title)
        self._cover_page(logo, [
            (214.5, [C(t, 22) for t in m["title_th"] + m["title_en"]]),
            (463.4, [C(m["author_th"], 20)]),
            (646.8, [C(t, 18) for t in foot])], first=True)
        # 2 Thai inner cover (+ student id, no English title)
        self._cover_page(logo, [
            (214.5, [C(t, 22) for t in m["title_th"]]),
            (463.4, [C(m["author_th"], 20), C(m["student_id"], 20)]),
            (673.1, [C(t, 18) for t in foot])], first=False)
        # 3 English inner cover - all capitals
        self._cover_page(logo, [
            (214.5, [C(t.upper(), 22) for t in m["title_en_upper"]]),
            (432.0, [C(m["author_en"].upper(), 20)]),
            (586.6, [C(t, 18) for t in (m.get("project_type_en", "A COMPUTER PROJECT"),
                                        "SUBMITTED IN PARTIAL FULFILLMENT OF",
                                        "THE REQUIREMENTS FOR THE DEGREE OF",
                                        f"BACHELOR OF SCIENCE IN {m['program_en'].upper()}")]
             + [(None, 18, False, "left", 0), C(m["university_en"].upper(), 18), C(str(m["year_ce"]), 18)])],
            first=False)
        # 4 approval page (16 pt regular, blank lines as in the approved book)
        B = (None, 16, False, "left", 0)
        N = lambda t: (t, 16, False, "center", 0)
        sig = lambda t: (t, 16, False, "center", 4320)
        L = lambda t: (t, 16, False, "left", 0)
        signers = [(m.get("advisor_label", "อาจารย์ที่ปรึกษา"), m["advisor"])]
        signers += [("กรรมการสอบ", name) for name in m.get("examiners", [])]
        signers += [(f"ประธานหลักสูตร{m['program_th']}", m["chair"])]
        lines = [N(m["university_th"]), N(m["faculty_th"]), B,
                 N(m.get("project_type_th", "โครงงานคอมพิวเตอร์")), N(m["author_prefix_th"] + m["author_th"]), B,
                 N("เรื่อง"), B] + [N(t) for t in m["title_th"]] + [B,
                 N("ได้รับการตรวจสอบและอนุมัติ ให้เป็นส่วนหนึ่งของการศึกษาตามหลักสูตร"), N(degree_line),
                 N(m.get("exam_date_line", f"เมื่อ วันที่ {m.get('exam_date_th', '')}"))]
        for label, name in signers:
            lines += [B, ((label, "__________________"), 16, False, "labelline", 4900), sig(f"({name})")]
        self._cover_page(logo, [(214.6, lines)], first=False)

    def abstract(self):
        m = self.data["meta"]; a = self.data["abstract"]
        self.new_section(number="roman", restart=1, hide_first=False)
        rows = [("ชื่อเรื่อง", " ".join(m["title_th"])), ("ชื่อนักศึกษา", m["author_prefix_th"] + m["author_th"]),
                ("ชื่อปริญญา", f"{m['degree_th']} สาขาวิชา{m['program_th']}"),
                ("อาจารย์ที่ปรึกษา", m["advisor"]), ("ปีพ.ศ.", str(m["year_be"]))]
        for lab, val in rows:
            par = self.p("", align="left", left=2160, hanging=2160, tabs=[("left", 2160, None)])
            add_run(par, lab, bold=True); add_tab(par); add_run(par, val)
        self.blank()
        self.p("บทคัดย่อ", align="left", size=18, bold=True)
        for t in a["paragraphs"]:
            self.p(t, first=720)
        self.blank()
        par = self.p("", align="thai")
        add_run(par, "คำสำคัญ:", bold=True); add_run(par, " " + " ".join(thai_sorted(a["keywords"])))
        # acknowledgement
        par = self.p("กิตติกรรมประกาศ", align="center", size=18, bold=True, page_break=True); self.blank()
        for t in self.data["acknowledgement"]:
            self.p(t, first=720)
        self.blank()
        self.p(m["author_th"], align="right")

    def _toc_title(self, title, page_break=True):
        self.p(title, align="center", size=20, bold=True, page_break=page_break, keep_next=True)
        self.blank()

    def lists_placeholders(self):
        """Write TOC, list of tables, list of figures with placeholder page
        numbers ('0'). finalize_thesis.py fills the numbers and inserts the
        '(ต่อ)' continuation pages."""
        tabs_dot = [("right", TEXT_W, "dot")]
        tabs_none = [("right", TEXT_W, None)]
        self._toc_title("สารบัญ")
        hdr = self.p("", align="left", tabs=tabs_none); add_tab(hdr, bold=True); add_run(hdr, "หน้า", bold=True)
        hdr._p.set(qn("w14:paraId") if False else "tocrole", "header")
        for kind, text, level in self.toc:
            if kind == "gap":
                self.blank(); continue
            if kind == "chapter":     # chapter lines: bold, no dots, bold number
                par = self.p("", align="left", tabs=tabs_none)
                add_run(par, text, bold=True); add_tab(par); add_run(par, "0", bold=True)
            elif kind == "main":      # front matter / bibliography / appendix / bio
                par = self.p("", align="left", tabs=tabs_dot)
                add_run(par, text, bold=True); add_tab(par, bold=False); add_run(par, "0", bold=True)
            else:                     # x.y sections, appendix items
                par = self.p("", align="left", left=720, tabs=tabs_dot)
                add_run(par, text); add_tab(par); add_run(par, "0")
            par._p.set("tocrole", kind)
        for title, label, items in (("สารบัญตาราง", "ตารางที่", self.lot), ("สารบัญภาพ", "ภาพที่", self.lof)):
            self._toc_title(title)
            hdr = self.p("", align="left", tabs=tabs_none)
            add_run(hdr, label, bold=True); add_tab(hdr); add_run(hdr, "หน้า", bold=True)
            for num, cap in items:
                par = self.p("", align="left", left=720, hanging=720, tabs=tabs_dot)
                add_run(par, f"{num} {cap}"); add_tab(par); add_run(par, "0")

    # ---------------------------------------------------------------- chapters
    def chapter(self, ch, first):
        self.new_section(number="arabic", restart=1 if first else None)
        n = ch["number"]
        self.p(f"บทที่ {n}", align="center", size=20, bold=True, keep_next=True, outline=0)
        self.p(ch["title"], align="center", size=20, bold=True, keep_next=True)
        self.blank()
        self.toc.append(("gap", None, 0))
        self.toc.append(("chapter", f"บทที่ {n} {ch['title']}", 0))
        self.fig[str(n)] = 0; self.tab[str(n)] = 0
        self.blocks(ch["blocks"], str(n))

    def blocks(self, blocks, chap, after_heading=True):
        level = 2          # body indent context: 2 -> first line 720, 3 -> 1440
        prev = "start"
        for b in blocks:
            t = b["type"]
            if t == "h2":
                if prev != "start":
                    self.blank()
                self.p(b["text"], align="left", size=18, bold=True, keep_next=True, outline=1)
                self.toc.append(("section", b["text"], 1)); level = 2
            elif t == "h3":
                self.p(b["text"], align="left", first=720, bold=True, keep_next=True, outline=2)
                level = 3
            elif t == "p":
                first = b.get("first", 1440 if level == 3 else 720)
                self.p(b["text"], first=first, align=body_align(b["text"]))
            elif t == "list":
                first = 1440 if level == 3 else 720
                for i, item in enumerate(b["items"], 1):
                    mark = f"{i}. " if b.get("style", "number") == "number" else "- "
                    self.p(mark + item, first=first + b.get("indent", 0), align=body_align(item))
            elif t == "item":     # labelled item: bold lead-in + regular description
                first = 1440 if level == 3 else 720
                par = self.p("", first=first)
                add_run(par, b["label"], bold=True)
                if b.get("text"): add_run(par, " " + b["text"])
            elif t == "figure":
                self.figure(b, chap)
            elif t == "table":
                self.table(b, chap)
            elif t == "equation":
                self.equation(b)
            elif t == "where":    # as in the approved book: "เมื่อ" regular at the margin,
                n = len(b["items"])   # items aligned at 0.5", one blank line after the block
                for i, (sym, desc) in enumerate(b["items"]):
                    par = self.p("", align="left", left=720, hanging=720 if i == 0 else 0,
                                 tabs=[("left", 720, None)], keep_next=i < n - 1)
                    if i == 0: add_run(par, "เมื่อ"); add_tab(par)
                    add_run(par, f"{sym} คือ {desc}")
                self.blank()
            elif t == "blank":
                self.blank()
            prev = t if t != "blank" else "blank"

    def figure(self, b, chap):
        self.fig[chap] = self.fig.get(chap, 0) + 1
        num = f"{chap}.{self.fig[chap]}"
        self.blank(keep_next=True)
        self.image(self.path(b["image"]), width_cm=b.get("width_cm"), height_cm=b.get("height_cm"))
        par = self.p("", align="center", keep_next=bool(b.get("source")))
        add_run(par, f"ภาพที่ {num}", bold=True); add_run(par, " " + b["caption"])
        if b.get("source"):
            self.p(f"ที่มา : {b['source']}", align="center")
        self.blank()
        self.lof.append((num, b["caption"]))

    def table(self, b, chap):
        self.tab[chap] = self.tab.get(chap, 0) + 1
        num = f"{chap}.{self.tab[chap]}"
        self.blank(keep_next=True)
        cap = self.p("", align="left", keep_next=True)
        add_run(cap, f"ตารางที่ {num}", bold=True); add_run(cap, " " + b["caption"])
        rows = b["rows"]; header = b.get("header_rows", [b["header"]] if b.get("header") else [])
        allrows = header + rows
        ncol = max(len(r) for r in allrows)
        tbl = self.d.add_table(rows=len(allrows), cols=ncol)
        widths = b.get("widths") or [TEXT_W // ncol] * ncol
        scale = TEXT_W / sum(widths) if sum(widths) > TEXT_W else 1
        widths = [int(w * scale) for w in widths]
        t = tbl._tbl
        tblPr = t.tblPr
        for e in list(tblPr):
            if e.tag in (qn("w:tblStyle"),):
                tblPr.remove(e)
        tw = OxmlElement("w:tblW"); tw.set(qn("w:w"), str(sum(widths))); tw.set(qn("w:type"), "dxa")
        old = tblPr.find(qn("w:tblW"))
        if old is not None: tblPr.remove(old)
        tblPr.append(tw)
        borders = OxmlElement("w:tblBorders")
        sides = ("top", "left", "bottom", "right", "insideH", "insideV") if b.get("grid") else ("top", "bottom", "insideH", "insideV")
        for s in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = OxmlElement("w:" + s)
            if s in sides:
                e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:space"), "0"); e.set(qn("w:color"), "000000")
            else:
                e.set(qn("w:val"), "nil")
            borders.append(e)
        tblPr.append(borders)
        mar = OxmlElement("w:tblCellMar")
        for s, v in (("top", 40), ("left", 80), ("bottom", 40), ("right", 80)):
            e = OxmlElement("w:" + s); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); mar.append(e)
        tblPr.append(mar)
        lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
        grid = t.find(qn("w:tblGrid"))
        for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
            gc.set(qn("w:w"), str(w))
        aligns = b.get("align") or (["left"] + ["center"] * (ncol - 1))
        nh = len(header)
        for ri, row in enumerate(allrows):
            tr = tbl.rows[ri]
            trPr = tr._tr.get_or_add_trPr()
            trPr.append(OxmlElement("w:cantSplit"))
            if ri < nh: trPr.append(OxmlElement("w:tblHeader"))
            for ci in range(ncol):
                cell = tr.cells[ci]; cell.width = Emu(widths[ci] * 635)
                tcPr = cell._tc.get_or_add_tcPr()
                va = OxmlElement("w:vAlign"); va.set(qn("w:val"), "center"); tcPr.append(va)
                val = row[ci] if ci < len(row) else ""
                par = cell.paragraphs[0]
                bold = ri < nh or (isinstance(val, str) and val.startswith("**"))
                if isinstance(val, str) and val.startswith("**"): val = val.strip("*")
                al = "center" if ri < nh else aligns[ci]
                fmt(par, align=al)
                if val == "<->":   # arrow spans the whole merged range if the cell is merged
                    span = [m for m in b.get("merge", []) if m[0] <= ri <= m[2] and m[1] == ci]
                    self._arrow(par, sum(widths[ci:span[0][3] + 1]) if span else widths[ci])
                elif val:
                    add_marked(par, str(val), bold=bold)
        for (r1, c1, r2, c2) in b.get("merge", []):
            a = tbl.cell(r1, c1); z = tbl.cell(r2, c2)
            keep = a.text
            m = a.merge(z)
            for extra in m.paragraphs[1:]:
                extra._p.getparent().remove(extra._p)
        # keep the whole table on one page (unless it is longer than a page):
        # every row's paragraphs keep-with-next, the last row does not
        if b.get("keep_together", True):
            # walk the raw <w:p> elements: tr.cells skips the lower part of a
            # vertically merged cell, which then breaks the keep-with-next chain
            for tr in tbl.rows[:-1]:
                for cp_el in tr._tr.iter(qn("w:p")):
                        pPr = cp_el.get_or_add_pPr()
                        if pPr.find(qn("w:keepNext")) is None:
                            pPr.insert(0, OxmlElement("w:keepNext"))
        # move table right after caption paragraph
        cap._p.addnext(t)
        if b.get("source"):
            src = self.p(f"ที่มา : {b['source']}", align="left")
        self.blank()
        self.lot.append((num, b["caption"]))

    def _arrow(self, par, width_twips):
        """Double-headed horizontal arrow drawn as a line shape centred in
        the cell - TH SarabunPSK has no arrow glyphs, typing '↔' would fall
        back to another font."""
        inner = max(width_twips - 160, 100)          # minus left/right cell margins (80 + 80)
        w = int(inner * 635 * 0.8)
        off = int(inner * 635 * 0.1)                 # offset is measured from the text area
        xml = f'''<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
 xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
 xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
 xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
<w:drawing><wp:anchor distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="2" behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1">
<wp:simplePos x="0" y="0"/><wp:positionH relativeFrom="column"><wp:posOffset>{off}</wp:posOffset></wp:positionH>
<wp:positionV relativeFrom="paragraph"><wp:posOffset>120000</wp:posOffset></wp:positionV>
<wp:extent cx="{w}" cy="0"/><wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapNone/><wp:docPr id="{id(par) % 100000 + 1000}" name="Arrow"/><wp:cNvGraphicFramePr/>
<a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"><wps:wsp><wps:cNvCnPr/>
<wps:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{w}" cy="0"/></a:xfrm><a:prstGeom prst="straightConnector1"><a:avLst/></a:prstGeom>
<a:ln w="15875"><a:solidFill><a:srgbClr val="000000"/></a:solidFill><a:headEnd type="arrow"/><a:tailEnd type="arrow"/></a:ln></wps:spPr><wps:bodyPr/></wps:wsp>
</a:graphicData></a:graphic></wp:anchor></w:drawing></w:r>'''
        from lxml import etree
        par._p.append(etree.fromstring(xml))

    def equation(self, b):
        if b.get("image"):   # pre-rendered equation: x-bar, sigma and radicals are not in TH SarabunPSK
            self.blank(keep_next=True)
            self.image(self.path(b["image"]), width_cm=b.get("width_cm"), height_cm=b.get("height_cm"))
            self.blank()
            return
        self.blank()
        par = self.p("", align="center", keep_next=True)
        omp = mel("oMathPara"); om = mel("oMath"); omp.append(om)
        def mr(text):
            r = mel("r"); rp = mel("rPr"); rp.append(mel("nor")); r.append(rp)
            w = OxmlElement("w:rPr"); set_rfonts(w, "Times New Roman")
            w.find(qn("w:rFonts")).set(qn("w:cs"), FONT)   # Thai words inside an equation
            for tag in ("w:sz", "w:szCs"):
                e = OxmlElement(tag); e.set(qn("w:val"), "22"); w.append(e)
            if THAI.search(text):
                w.append(OxmlElement("w:cs"))
                lg = OxmlElement("w:lang"); lg.set(qn("w:bidi"), "th-TH"); w.append(lg)
            r.append(w); t = mel("t"); t.set(XML_SPACE, "preserve"); t.text = text; r.append(t); return r
        if b.get("num"):
            om.append(mr(b.get("lhs", "") + " = " if b.get("lhs") else ""))
            f = mel("f"); n = mel("num"); dn = mel("den")
            n.append(mr(b["num"])); dn.append(mr(b["den"])); f.append(n); f.append(dn)
            if b.get("sqrt"):   # square root over the whole fraction (standard deviation)
                rad = mel("rad"); rp = mel("radPr"); dh = mel("degHide"); dh.set(f"{{{M}}}val", "1")
                rp.append(dh); rad.append(rp); rad.append(mel("deg")); e = mel("e"); e.append(f); rad.append(e)
                om.append(rad)
            else:
                om.append(f)
        else:
            om.append(mr(b["text"]))
        par._p.append(omp)
        after = self.blank()
        if after is not None:   # keep the equation with its "เมื่อ" block / next line
            fmt(after, align="left", keep_next=True)

    # -------------------------------------------------------------- back matter
    def references(self):
        refs = self.data["references"]
        self.new_section(number="arabic", hide_first=True)
        self.p("บรรณานุกรม", align="center", size=20, bold=True, keep_next=True); self.blank()
        self.toc.append(("gap", None, 0)); self.toc.append(("main", "บรรณานุกรม", 0))
        for e in thai_sorted(refs.get("thai", [])):
            par = self.p("", align="thai", left=720, hanging=720)
            add_marked(par, url_breaks(e))
        for e in sorted(refs.get("english", []), key=apa_key):
            par = self.p("", align="both", left=720, hanging=720)
            add_marked(par, url_breaks(e))

    def appendices(self):
        apps = self.data.get("appendices", [])
        if not apps: return
        self.new_section(number="arabic", hide_first=True)
        self.spacer(364.1 - 108)
        self.p("ภาคผนวก", align="center", size=18, bold=True)
        self.toc.append(("main", "ภาคผนวก", 0))
        for ap in apps:
            L = ap["letter"]
            self.new_section(number="arabic", hide_first=True)
            self.spacer(419.0 - 108)
            self.p(f"ภาคผนวก {L}", align="center", size=18, bold=True)
            self.p(ap["title"], align="center", size=18, bold=True)
            self.toc.append(("section", f"ภาคผนวก {L} {ap['title']}", 1))
            self.new_section(number="arabic", hide_first=True)
            self.p(f"ภาคผนวก {L}", align="center", size=18, bold=True)
            self.p(ap["title"], align="center", size=18, bold=True); self.blank()
            self.blocks(ap["blocks"], L)

    def bio(self):
        b = self.data["bio"]
        self.new_section(number="arabic", hide_first=False)
        self.p("ประวัติผู้จัดทำ", align="center", size=20, bold=True); self.blank()
        self.toc.append(("main", "ประวัติผู้จัดทำ", 0))
        if b.get("photo"):
            self.image(self.path(b["photo"]), width_cm=3.0); self.blank()
        for lab, val in b["rows"]:
            par = self.p("", align="left", left=2160, hanging=2160, tabs=[("left", 2160, None)])
            add_run(par, lab, bold=True); add_tab(par); add_run(par, val)

    # ----------------------------------------------------------------- build
    def build(self, out):
        self.covers()
        self.abstract()
        # TOC entries for front matter are known up front
        self.toc[:0] = [("main", "บทคัดย่อ", 0), ("main", "กิตติกรรมประกาศ", 0), ("main", "สารบัญ", 0),
                        ("main", "สารบัญตาราง", 0), ("main", "สารบัญภาพ", 0)]
        marker = self.d.add_paragraph(); marker._p.set("lists", "here")
        for i, ch in enumerate(self.data["chapters"]):
            self.chapter(ch, i == 0)
        self.references()
        self.appendices()
        self.bio()
        # now we know every caption/heading -> write the lists before the marker
        body = self.d.element.body
        before = len(body)
        self.lists_placeholders()
        new = list(body)[before - 1:]
        # lists_placeholders appended at the end (inside last section); move them
        sect = body.find(qn("w:sectPr"))
        moved = [e for e in new if e is not sect]
        for e in moved:
            marker._p.addprevious(e)
        body.remove(marker._p)
        for e in body.iter():
            for a in ("tocrole", "lists"):
                if a in e.attrib: pass
        self.d.save(out)


def body_align(text):
    """thaiDistribute letter-spaces Thai words when a line holds long English
    tokens (Next.js Route Handlers ...); justify on spaces instead for such text."""
    letters = [c for c in text if c.isalpha()]
    latin = sum(1 for c in letters if c.isascii())
    return "both" if letters and latin / len(letters) > 0.15 else "thai"


def thai_sorted(items):
    """Dictionary order for Thai: leading vowels เ แ โ ใ ไ are skipped
    when comparing (ฉ before เฉ...), then English a-z."""
    def key(s):
        s2 = re.sub(r"^([เแโใไ])(.)", r"\2\1", s.strip())
        return s2
    th = [x for x in items if THAI.match(x.strip())]
    en = [x for x in items if not THAI.match(x.strip())]
    return sorted(th, key=key) + sorted(en, key=str.lower)


def apa_key(s):
    """APA order: by author names letter by letter, ignoring punctuation and
    '&' (so 'Jocher, G., Chaurasia, A., & Qiu' precedes 'Jocher, G., & Qiu')."""
    s = re.sub(r"[*&.,]", " ", s.lower())
    return re.sub(r"\s+", " ", s).strip()


def url_breaks(text):
    """Insert zero-width spaces after / . - _ in URLs so long links can wrap
    instead of pushing the whole link to the next line (which makes the
    previous justified line stretch badly)."""
    zw = "​"
    def brk(m):
        u = m.group(0); head = re.match(r"https?://", u)
        pre = head.group(0) if head else ""
        return pre + zw + re.sub(r"([/._\-?=&])", lambda k: k.group(1) + zw, u[len(pre):])
    return re.sub(r"https?://[^\s]+", brk, text)


def main():
    src, out = sys.argv[1], sys.argv[2]
    data = json.load(open(src, encoding="utf-8"))
    b = Builder(data, os.path.dirname(os.path.abspath(src)))
    b.build(out)
    # strip helper attributes
    d = docx.Document(out)
    for e in d.element.body.iter():
        for a in ("tocrole", "lists"):
            if a in e.attrib: del e.attrib[a]
    d.save(out)
    print("built", out)


if __name__ == "__main__":
    main()
